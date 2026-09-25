# -*- coding: utf-8 -*-
"""JANスキャン棚卸し(Step 1 MVP)のテスト

評価観点(2026-09-25 指示)ごとにテストを分けている:
    A. JAN から正しい P番号へ特定できるか(旧システムのIDを使わない・P番号の一意性)
    B. 未登録JANでも作業が止まらないか
    C. 同一商品を複数回数えても正しく合算されるか
    D. 入力ミスを検出できるか
    E. 生データから集計を再現できるか(追記のみ・取消イベント・1件ごとの永続化)
    F. セット梱包品の将来換算(20 + 10×2 = 40)に拡張できるか
    G. 画面(HTTP)経由でも同じ規則で動き、localhost/LAN の保護が効くか

ダミーの商品マスターは一時フォルダに作る。実マスターは読み取り専用で開き、ハッシュが変わらないことを確かめる。

実行: python3 -m unittest discover -s MomijiStore_OS/01_InventoryManagement/tests -v
"""
import csv
import json
import os
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from datetime import datetime, timedelta
from http.server import ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'Python' / 'Stocktake'))
import jan_lookup as J  # noqa: E402
import stocktake_aggregate as A  # noqa: E402
import stocktake_app as APP  # noqa: E402
import stocktake_store as S  # noqa: E402

CFG = J.load_config()


def ean(body):
    """チェックディジットを付けた正しいJANを作る。"""
    for d in '0123456789':
        if J.gtin_check_ok(body + d):
            return body + d


GOAT = '4514610005035'           # 実在: P000056 ゴートミルク
TOWEL = '4533141407514'          # 実在: 入数違いの3商品が同じJAN
SET_JAN = ean('490000000001')    # セットにだけ付いたJAN
SHARED = ean('490000000002')     # 単品とセットで共通
NEW_JAN = ean('490000000009')    # マスター未登録
UPC = '192333006122'             # UPC-A(12桁)。0を補うと 0192333006122

MASTER_HEADER = ['内部管理ID', 'JAN', '商品名', '商品種別', '標準原価', '販売チャネル']
BASE_ROWS = [
    ['P000001', GOAT, 'ゴートミルク 120g', '単品', 1270, '両方'],
    ['P000002', TOWEL, 'ネックタオル 50本入り', '単品', 500, '楽天'],
    ['P000003', TOWEL, 'ネックタオル 100本入り', '単品', 900, '楽天'],
    ['P000004', SET_JAN, 'ゴートミルク 2袋セット', 'セット', 2540, 'Amazon'],
    ['P000005', SHARED, '共通JANの単品', '単品', 300, '楽天'],
    ['P000006', SHARED, '共通JANの3個セット', 'セット', 900, 'Amazon'],
    ['P000007', None, 'JANなしの単品', '単品', 100, '楽天'],
    ['P000008', UPC, 'UPCの単品', '単品', None, '楽天'],
    ['P000009', 'B0DJLXVQSH', 'JAN欄にASIN', '単品', 3940, 'Amazon'],
]


def make_master(folder, rows=None, header=None, name='商品マスター_単品_v1.0.xlsx'):
    wb = Workbook()
    ws = wb.active
    ws.title = J.SH_MASTER
    ws.append(header or MASTER_HEADER)
    for r in rows if rows is not None else BASE_ROWS:
        ws.append(r)
    path = os.path.join(folder, name)
    wb.save(path)
    return path


class Clock:
    def __init__(self):
        self.t = datetime(2026, 10, 1, 9, 0, 0)

    def __call__(self):
        return self.t

    def tick(self, sec=30):
        self.t += timedelta(seconds=sec)


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = self.tmp.name
        self.master = J.load_master(make_master(self.dir))
        self.clock = Clock()
        self.store = S.Store(self.master, CFG, data_dir=os.path.join(self.dir, 'data'), clock=self.clock)
        self.sid = self.store.start_session('英樹', 'MAC-1')

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def reg(self, raw, qty, loc='', device='MAC-1', **kw):
        self.clock.tick()
        return self.store.register(self.sid, '英樹', device, loc, raw, qty, **kw)

    def agg(self, comp=None, store=None):
        return A.run([self.sid], store or self.store, set_composition=comp or {})

    def product(self, result, pid):
        return next((b for b in result['products'] if b['pid'] == pid), None)


# ─────────────── A. JAN → 正しい P番号 ───────────────

class NormalizeCode(unittest.TestCase):
    def n(self, raw):
        return J.normalize_code(raw, CFG)

    def test_ean13(self):
        self.assertEqual(self.n(GOAT), {'kind': J.CODE_JAN, 'code': GOAT, 'message': ''})

    def test_fullwidth_and_spaces(self):
        # iPad の日本語入力で全角になっても同じJAN
        self.assertEqual(self.n(' ４５１４６１０００５０３５ ')['code'], GOAT)

    def test_upc_a_gets_leading_zero(self):
        self.assertEqual(self.n(UPC)['code'], '0' + UPC)

    def test_gtin14_indicator0_is_same_item(self):
        self.assertEqual(self.n('0' + GOAT)['code'], GOAT)

    def test_case_code_rejected(self):
        body = '1' + GOAT[:-1]
        self.assertEqual(self.n(ean(body))['kind'], J.CODE_CASE)
        self.assertEqual(self.n('1234567890123456')['kind'], J.CODE_CASE)

    def test_bad_check_digit(self):
        bad = GOAT[:-1] + str((int(GOAT[-1]) + 1) % 10)
        self.assertEqual(self.n(bad)['kind'], J.CODE_INVALID)

    def test_ean8(self):
        self.assertEqual(self.n(ean('4901234'))['kind'], J.CODE_JAN)

    def test_location_temp_command(self):
        self.assertEqual(self.n('a-01-02')['kind'], J.CODE_LOCATION)
        self.assertEqual(self.n('T-000001')['kind'], J.CODE_TEMP)
        self.assertEqual(self.n('CMD-UNDO')['kind'], J.CODE_COMMAND)

    def test_master_key(self):
        self.assertEqual(J.master_jan_key(UPC)[0], '0' + UPC)
        self.assertEqual(J.master_jan_key('B0DJLXVQSH'), (None, 'ASINが入っている'))
        self.assertEqual(J.master_jan_key(None), (None, '空欄'))
        self.assertEqual(J.master_jan_key(4514610005035.0)[0], GOAT)


class MasterIndex(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.m = J.load_master(make_master(self.tmp.name))

    def tearDown(self):
        self.tmp.cleanup()

    def test_match(self):
        self.assertEqual(self.m.resolve_jan(GOAT), {'status': J.R_MATCH, 'pid': 'P000001', 'candidates': ['P000001']})

    def test_duplicate_is_not_guessed(self):
        r = self.m.resolve_jan(TOWEL)
        self.assertEqual((r['status'], r['pid'], r['candidates']), (J.R_DUPLICATE, '', ['P000002', 'P000003']))

    def test_shared_jan_is_not_auto_confirmed(self):
        # 必須修正①: 単品とセットが同じJANなら単品へ自動確定しない
        self.assertEqual(self.m.resolve_jan(SHARED),
                         {'status': J.R_SHARED, 'pid': '', 'candidates': ['P000005', 'P000006']})

    def test_set_only(self):
        self.assertEqual(self.m.resolve_jan(SET_JAN)['status'], J.R_SET_ONLY)

    def test_upc_in_master_matches_padded_scan(self):
        self.assertEqual(self.m.resolve_jan('0' + UPC)['pid'], 'P000008')

    def test_asin_in_jan_column_is_reported_not_indexed(self):
        self.assertIn(('P000009', 'B0DJLXVQSH', 'ASINが入っている'), self.m.jan_issues)

    def test_duplicate_pid_stops(self):
        rows = BASE_ROWS + [['P000001', ean('490000000003'), '別商品', '単品', 1, '楽天']]
        with self.assertRaisesRegex(J.MasterError, 'P000001 が重複'):
            J.load_master(make_master(self.tmp.name, rows, name='dup.xlsx'))

    def test_bad_pid_format_stops(self):
        with self.assertRaisesRegex(J.MasterError, '形式が不正'):
            J.load_master(make_master(self.tmp.name, BASE_ROWS + [['P12', GOAT, 'x', '単品', 1, '']], name='bad.xlsx'))

    def test_legacy_system_layout_is_refused(self):
        # 旧 在庫管理システムの見出し(🔑 内部管理ID / JANコード)では読み込ませない
        legacy = ['🔑 内部管理ID', '商品名', 'JANコード', 'ASIN']
        with self.assertRaises(J.MasterError):
            J.load_master(make_master(self.tmp.name, [['P000001', 'ロイヤルカナン', None, 'B0DJLXVQSH']],
                                      header=legacy, name='legacy_like.xlsx'))

    def test_legacy_file_name_is_refused(self):
        p = make_master(self.tmp.name, name='在庫管理システム_v1.0.xlsm')
        with self.assertRaisesRegex(J.MasterError, 'ID体系が異なる'):
            J.load_master(p)


@unittest.skipUnless(os.path.exists(J.MASTER_FILE), '実マスターが無い環境')
class RealMasterReadOnly(unittest.TestCase):
    """実データで照合できること。正本を1バイトも変えないこと。"""

    @classmethod
    def setUpClass(cls):
        cls.before = {p: J.sha256_file(p) for p in (J.MASTER_FILE, J.LEGACY_SYSTEM_FILE) if os.path.exists(p)}
        cls.m = J.load_master()

    def test_default_master_is_the_official_one(self):
        self.assertEqual(os.path.basename(J.MASTER_FILE), '商品マスター_単品_v1.0.xlsx')
        self.assertIn('SourceData', J.MASTER_FILE)

    def test_goat_milk_is_p000056(self):
        r = self.m.resolve_jan(GOAT)
        self.assertEqual((r['status'], r['pid']), (J.R_MATCH, 'P000056'))
        self.assertEqual(self.m.products['P000056']['kind'], '単品')

    def test_real_duplicate_returns_candidates(self):
        r = self.m.resolve_jan(TOWEL)
        self.assertEqual(r['status'], J.R_DUPLICATE)
        self.assertTrue({'P000555', 'P000556', 'P000557'} <= set(r['candidates']))

    def test_legacy_real_file_refused(self):
        if os.path.exists(J.LEGACY_SYSTEM_FILE):
            with self.assertRaises(J.MasterError):
                J.load_master(J.LEGACY_SYSTEM_FILE)

    def test_files_unchanged(self):
        for p, h in self.before.items():
            self.assertEqual(J.sha256_file(p), h, p)


# ─────────────── B. 未登録JANでも止まらない ───────────────

class Unregistered(Base):
    def test_unregistered_is_recorded_without_search(self):
        e = self.reg(NEW_JAN, 7, loc='B-02-01')
        self.assertEqual((e['照合結果'], e['内部管理ID'], e['数量'], e['数量単位']), (J.R_UNREGISTERED, '', '7', S.UNIT_PIECE))
        r = self.agg()
        self.assertEqual([(b['jan'], b['qty'], dict(b['locs'])) for b in r['unregistered']], [(NEW_JAN, 7, {'B-02-01': 7})])
        self.assertEqual(r['stats']['units'], 0)

    def test_resolved_automatically_after_master_registration(self):
        self.reg(NEW_JAN, 7)
        raw_hash = J.sha256_file(self.store.events_path(self.sid))
        rows = BASE_ROWS + [['P000010', NEW_JAN, '後で登録した商品', '単品', 400, '楽天']]
        self.store.master = J.load_master(make_master(self.dir, rows, name='m2.xlsx'))
        r = self.agg()
        self.assertEqual(self.product(r, 'P000010')['qty'], 7)
        self.assertEqual(r['unregistered'], [])
        self.assertTrue(any(c[0] == A.CHK_RESOLVED_LATER for c in r['checks']))
        self.assertEqual(J.sha256_file(self.store.events_path(self.sid)), raw_hash)   # 生データは書き換えない

    def test_jan_less_item_gets_temp_id_and_can_be_linked(self):
        e1 = self.store.register_temp(self.sid, '英樹', 'MAC-1', '', '青い箱 犬用ガム', 4)
        self.clock.tick()
        e2 = self.store.register_temp(self.sid, '英樹', 'MAC-1', '', '赤い袋', 2)
        self.assertEqual((e1['正規化コード'], e2['正規化コード']), ('T-000001', 'T-000002'))
        self.clock.tick()
        self.reg('t-000001', 3)                                   # 仮IDラベルを後でスキャン
        r = self.agg()
        self.assertEqual({b['temp']: b['qty'] for b in r['temps']}, {'T-000001': 7, 'T-000002': 2})
        with open(os.path.join(self.store.dir, S.TEMP_LINKS_FILE), 'a', encoding='utf-8', newline='') as f:
            csv.writer(f).writerow(['T-000001', 'P000007', '2026-10-02 10:00:00', '英樹', '現物確認'])
        r = self.agg()
        self.assertEqual(self.product(r, 'P000007')['qty'], 7)
        self.assertEqual([b['temp'] for b in r['temps']], ['T-000002'])

    def test_temp_requires_memo_and_valid_qty_before_issuing(self):
        with self.assertRaises(S.StoreError):
            self.store.register_temp(self.sid, '英樹', 'MAC-1', '', '', 1)
        with self.assertRaises(S.StoreError):
            self.store.register_temp(self.sid, '英樹', 'MAC-1', '', 'メモ', '')
        self.assertEqual(self.store.temp_items(), {})             # 孤立した仮IDを作らない

    def test_unknown_temp_label_rejected(self):
        with self.assertRaisesRegex(S.StoreError, '発行されていません'):
            self.reg('T-000099', 1)


# ─────────────── C. 複数回の合算 ───────────────

class Summation(Base):
    def test_two_shelves_sum_to_50(self):
        self.reg(GOAT, 30, loc='A-01-01')
        self.reg(GOAT, 20, loc='C-01-01')
        b = self.product(self.agg(), 'P000001')
        self.assertEqual((b['qty'], b['count'], dict(b['locs'])), (50, 2, {'A-01-01': 30, 'C-01-01': 20}))

    def test_other_device_also_sums(self):
        self.reg(GOAT, 30, device='MAC-1')
        self.reg(GOAT, 5, device='IPAD-1')
        self.assertEqual(self.product(self.agg(), 'P000001')['qty'], 35)

    def test_same_place_same_qty_is_flagged_not_removed(self):
        self.reg(GOAT, 10, loc='A-01-01')
        self.clock.tick(60)
        self.reg(GOAT, 10, loc='A-01-01')
        r = self.agg()
        self.assertEqual(self.product(r, 'P000001')['qty'], 20)
        self.assertTrue(any(c[0] == A.CHK_REPEAT for c in r['checks']))

    def test_duplicate_jan_choice_and_unselected(self):
        self.reg(TOWEL, 5, chosen_pid='P000003')
        self.reg(TOWEL, 2)
        r = self.agg()
        self.assertEqual(self.product(r, 'P000003')['qty'], 5)
        self.assertIsNone(self.product(r, 'P000002'))
        self.assertEqual([(b['jan'], b['qty']) for b in r['unregistered']], [(TOWEL, 2)])
        self.assertTrue(any(c[0] == A.CHK_DUP_UNSELECTED for c in r['checks']))

    def test_zero_is_a_count(self):
        self.reg(GOAT, 0)
        b = self.product(self.agg(), 'P000001')
        self.assertEqual((b['qty'], b['count']), (0, 1))          # 0 と「数えていない」を区別する

    def test_amount_uses_standard_cost_and_flags_missing(self):
        self.reg(UPC, 3)
        r = self.agg()
        self.assertIsNone(self.product(r, 'P000008')['cost'])
        self.assertTrue(any(c[0] == A.CHK_NO_COST for c in r['checks']))


# ─────────────── D. 入力ミスの検出 ───────────────

class InputErrors(Base):
    def code(self, fn):
        with self.assertRaises(S.StoreError) as cm:
            fn()
        return cm.exception.code

    def test_quantity_rules(self):
        self.assertEqual(self.code(lambda: self.reg(GOAT, '')), 'qty_blank')
        self.assertEqual(self.code(lambda: self.reg(GOAT, '-1')), 'qty_invalid')
        self.assertEqual(self.code(lambda: self.reg(GOAT, '3.5')), 'qty_invalid')
        self.assertEqual(self.code(lambda: self.reg(GOAT, GOAT)), 'qty_looks_like_code')
        self.assertEqual(self.code(lambda: self.reg(GOAT, '10000')), 'qty_too_large')
        self.assertEqual(self.code(lambda: self.reg(GOAT, '300')), 'confirm_quantity')
        self.assertEqual(self.reg(GOAT, '300', qty_confirmed=True)['数量'], '300')
        self.assertEqual(self.reg(SHARED, '１２')['数量'], '12')   # 全角数字は受け付ける

    def test_errors_write_nothing(self):
        for q in ('', 'abc', GOAT):
            with self.assertRaises(S.StoreError):
                self.reg(GOAT, q)
        self.assertEqual(self.store.events(self.sid), [])

    def test_bad_codes_are_not_registered(self):
        bad = GOAT[:-1] + str((int(GOAT[-1]) + 1) % 10)
        self.assertEqual(self.code(lambda: self.reg(bad, 1)), 'bad_code')
        self.assertEqual(self.code(lambda: self.reg(ean('1' + GOAT[:-1]), 1)), 'bad_code')   # ケースコード
        self.assertEqual(self.code(lambda: self.reg('A-01-01', 1)), 'bad_code')             # ロケーションは商品ではない

    def test_duplicate_guard(self):
        self.reg(GOAT, 10, loc='A-01-01')
        self.clock.tick(5)
        with self.assertRaises(S.StoreError) as cm:
            self.store.register(self.sid, '英樹', 'MAC-1', 'A-01-01', GOAT, 10)
        self.assertEqual(cm.exception.code, 'confirm_duplicate')
        self.store.register(self.sid, '英樹', 'MAC-1', 'A-01-01', GOAT, 10, dup_confirmed=True)
        self.assertEqual(self.product(self.agg(), 'P000001')['qty'], 20)

    def test_bad_choice_rejected(self):
        self.assertEqual(self.code(lambda: self.reg(TOWEL, 1, chosen_pid='P000001')), 'bad_choice')
        self.assertEqual(self.code(lambda: self.reg(GOAT, 1, chosen_pid='P000002')), 'bad_choice')

    def test_location_format(self):
        self.assertEqual(self.code(lambda: self.reg(GOAT, 1, loc='棚の奥')), 'bad_location')

    def test_staff_required_and_formula_neutralised(self):
        self.assertEqual(self.code(lambda: self.store.register(self.sid, '', 'MAC-1', '', GOAT, 1)), 'no_staff')
        e = self.store.register(self.sid, '=HYPERLINK("x")', 'MAC-1', '', GOAT, 1)
        self.assertTrue(e['担当者'].startswith("'="))

    def test_closed_session_rejects(self):
        self.store.close_session(self.sid, '英樹', 'MAC-1')
        self.assertEqual(self.code(lambda: self.reg(GOAT, 1)), 'session_closed')


# ─────────────── E. 生データから集計を再現 ───────────────

class Reproducibility(Base):
    def fill(self):
        self.reg(GOAT, 30, loc='A-01-01')
        self.reg(GOAT, 20, loc='C-01-01')
        self.reg(NEW_JAN, 4)
        self.reg(TOWEL, 5, chosen_pid='P000002')
        self.reg(SET_JAN, 10)
        self.reg(SHARED, 8, device='IPAD-1', chosen_pid='P000005')
        self.store.register_temp(self.sid, '英樹', 'MAC-1', '', '青い箱', 2)

    def test_same_input_same_result(self):
        self.fill()
        r1, r2 = self.agg(), self.agg()
        self.assertEqual(r1['fingerprint'], r2['fingerprint'])
        self.store.close()
        reader = S.Store(self.master, CFG, data_dir=self.store.dir, writer=False)   # 別プロセス相当
        self.assertEqual(A.run([self.sid], reader, set_composition={})['fingerprint'], r1['fingerprint'])
        self.store = S.Store(self.master, CFG, data_dir=self.store.dir, clock=self.clock)

    def test_cancel_is_an_event_and_raw_rows_remain(self):
        self.fill()
        before = self.store.events(self.sid)
        self.clock.tick()
        c = self.store.cancel(self.sid, '英樹', 'MAC-1')          # MAC-1 の直前 = 仮ID
        self.assertEqual(c['イベント種別'], S.EV_CANCEL)
        self.clock.tick()
        c2 = self.store.cancel(self.sid, '英樹', 'MAC-1')         # 次は SET_JAN(IPAD-1 の登録は対象外)
        self.assertEqual(c2['正規化コード'], SET_JAN)
        after = self.store.events(self.sid)
        self.assertEqual(after[:len(before)], before)            # 既存行は1行も変わらない
        r = self.agg()
        self.assertEqual((r['stats']['cancelled'], r['temps'], r['set_pending']), (2, [], []))
        self.assertEqual(self.product(r, 'P000005')['qty'], 8)
        with self.assertRaises(S.StoreError):
            self.store.cancel(self.sid, '英樹', 'MAC-1', c2['取消対象イベントID'])

    def test_persisted_per_event_and_numbering_continues(self):
        self.reg(GOAT, 1)
        self.store.close()                                        # 強制終了相当
        self.store = S.Store(self.master, CFG, data_dir=os.path.join(self.dir, 'data'), clock=self.clock)
        e = self.reg(GOAT, 2)
        self.assertEqual(e['イベントID'], f'{self.sid}-00002')
        self.assertEqual(self.product(self.agg(), 'P000001')['qty'], 3)

    def test_second_writer_is_refused(self):
        with self.assertRaises(S.StoreError) as cm:
            S.Store(self.master, CFG, data_dir=self.store.dir)
        self.assertEqual(cm.exception.code, 'locked')

    def test_concurrent_devices(self):
        jans = [ean('49100000%04d' % i) for i in range(40)]
        errors = []

        def work(dev):
            try:
                for j in jans:
                    self.store.register(self.sid, '英樹', dev, '', j, 1)
            except Exception as ex:  # noqa: BLE001
                errors.append(ex)
        ts = [threading.Thread(target=work, args=(f'DEV-{i}',)) for i in range(5)]
        for t in ts:
            t.start()
        for t in ts:
            t.join()
        self.assertEqual(errors, [])
        evs = self.store.events(self.sid)
        self.assertEqual(len(evs), 200)
        self.assertEqual(len({e['イベントID'] for e in evs}), 200)
        self.assertEqual(sum(b['qty'] for b in self.agg()['unregistered']), 200)

    def test_xlsx_matches_result(self):
        self.fill()
        r = self.agg()
        out = A.write_xlsx(r, [self.sid], self.store, os.path.join(self.dir, 'out'))
        wb = load_workbook(out, read_only=True)
        self.assertEqual(wb.sheetnames, ['実棚集計', '未登録JAN', '仮ID', 'セット換算待ち', '要確認', '明細', '集計情報'])
        rows = {r[0]: r[3] for r in wb['実棚集計'].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows, {'P000001': 50, 'P000002': 5, 'P000005': 8})
        info = {r[0]: r[1] for r in wb['集計情報'].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(info['集計結果ハッシュ'], r['fingerprint'])
        self.assertEqual(sum(1 for _ in wb['明細'].iter_rows(min_row=2)), len(self.store.events(self.sid)))

    def test_session_ids(self):
        self.assertEqual(self.sid, 'ST-20261001-01')
        self.assertEqual(self.store.start_session('英樹', 'MAC-1'), 'ST-20261001-02')


# ─────────────── F. セット梱包品の換算(将来) ───────────────

class SetExtension(Base):
    def test_set_is_kept_apart_until_composition_exists(self):
        self.reg(GOAT, 20)
        e = self.reg(SET_JAN, 10)
        self.assertEqual((e['内部管理ID'], e['数量単位']), ('P000004', S.UNIT_SET))
        r = self.agg()
        self.assertEqual(self.product(r, 'P000001')['qty'], 20)                  # 単品に混ぜない
        self.assertEqual([(b['pid'], b['qty']) for b in r['set_pending']], [('P000004', 10)])

    def test_goat_milk_example_20_plus_10x2_is_40(self):
        self.reg(GOAT, 20)
        self.reg(SET_JAN, 10)
        b = self.product(self.agg(comp={'P000004': [('P000001', 2)]}), 'P000001')
        self.assertEqual((b['qty'], b['direct'], b['from_sets']), (40, 20, 20))

    def test_default_composition_is_empty_in_step1(self):
        self.assertEqual(A.load_set_composition(), {})


# ─────────────── H. 単品セット共通JAN(必須修正① 2026-09-25) ───────────────

class SharedJan(Base):
    """バラの単品 / 梱包済みセット / 後で確認 の3択。セットは換算せず「セット換算待ち」。"""

    def test_unselected_is_kept_for_review_not_counted(self):
        e = self.reg(SHARED, 8)
        self.assertEqual((e['照合結果'], e['内部管理ID'], e['数量単位']), (J.R_SHARED, '', S.UNIT_PIECE))
        self.assertIn('P000005', e['備考'])
        r = self.agg()
        self.assertIsNone(self.product(r, 'P000005'))                        # 単品に黙って入れない
        self.assertEqual([(b['jan'], b['qty'], b['reason']) for b in r['unregistered']],
                         [(SHARED, 8, A.CHK_SHARED_UNSELECTED)])
        self.assertTrue(any(c[0] == A.CHK_SHARED_UNSELECTED for c in r['checks']))

    def test_loose_single(self):
        e = self.reg(SHARED, 104, chosen_pid='P000005')
        self.assertEqual((e['照合結果'], e['内部管理ID'], e['数量単位']), (J.R_CHOSEN, 'P000005', S.UNIT_SINGLE))
        self.assertEqual(self.product(self.agg(), 'P000005')['qty'], 104)

    def test_packed_set_is_recorded_as_set_p_and_not_converted(self):
        e = self.reg(SHARED, 8, chosen_pid='P000006')
        self.assertEqual((e['内部管理ID'], e['数量'], e['数量単位']), ('P000006', '8', S.UNIT_SET))
        self.assertIn('セット換算待ち', e['備考'])
        r = self.agg()
        self.assertEqual([(b['pid'], b['qty']) for b in r['set_pending']], [('P000006', 8)])
        self.assertIsNone(self.product(r, 'P000005'))                        # 入数を推測して換算しない

    def test_rehearsal_case_loose_104_plus_packed_8(self):
        # リハーサルで 112(誤り)になったケース: バラ104 + 2本セット梱包8
        self.reg(SHARED, 104, chosen_pid='P000005')
        self.reg(SHARED, 8, chosen_pid='P000006')
        r = self.agg()
        self.assertEqual(self.product(r, 'P000005')['qty'], 104)
        self.assertEqual(r['set_pending'][0]['qty'], 8)
        # 差し込み口: 構成表が渡されたときだけ セット数×構成数量 を単品へ(ここはテスト用の仮の構成)
        b = self.product(self.agg(comp={'P000006': [('P000005', 3)]}), 'P000005')
        self.assertEqual((b['qty'], b['direct'], b['from_sets']), (128, 104, 24))

    def test_choice_must_be_a_candidate(self):
        with self.assertRaises(S.StoreError) as cm:
            self.reg(SHARED, 1, chosen_pid='P000001')
        self.assertEqual(cm.exception.code, 'bad_choice')

    def test_old_auto_confirmed_record_is_not_silently_counted(self):
        # 修正前の規則(単品へ自動確定)で記録された行も、再集計では要確認へ回る
        old_rows = [r for r in BASE_ROWS if r[0] != 'P000006']
        self.store.master = J.load_master(make_master(self.dir, old_rows, name='old.xlsx'))
        e = self.reg(SHARED, 8)
        self.assertEqual((e['照合結果'], e['内部管理ID']), (J.R_MATCH, 'P000005'))
        self.store.master = self.master
        r = self.agg()
        self.assertIsNone(self.product(r, 'P000005'))
        kinds = {c[0] for c in r['checks']}
        self.assertTrue({A.CHK_SHARED_UNSELECTED, A.CHK_CHANGED} <= kinds)


@unittest.skipUnless(os.path.exists(J.MASTER_FILE), '実マスターが無い環境')
class RealSharedJan21(unittest.TestCase):
    """実マスターの単品セット共通JAN 21件すべてを、3通り(バラ/セット/後で確認)で登録して集計する。"""

    @classmethod
    def setUpClass(cls):
        cls.m = J.load_master()
        cls.shared = sorted(set(cls.m.by_jan_single) & set(cls.m.by_jan_set))

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = S.Store(self.m, CFG, data_dir=self.tmp.name)
        self.sid = self.store.start_session('テスト', 'MAC-1')

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def test_count_is_21(self):
        # Step 0(2026-09-25)時点の件数。マスターが更新されて変わったら、このテストと手順書を見直す
        self.assertEqual(len(self.shared), 21, self.shared)

    def test_none_is_auto_confirmed(self):
        for jan in self.shared:
            r = self.m.resolve_jan(jan)
            kinds = {self.m.products[p]['kind'] for p in r['candidates']}
            self.assertEqual((r['status'], r['pid']), (J.R_SHARED, ''), jan)
            self.assertEqual(kinds, {J.KIND_SINGLE, J.KIND_SET}, jan)

    def test_three_ways_for_all_21(self):
        expect_single, expect_set = {}, {}
        for jan in self.shared:
            c = self.m.resolve_jan(jan)['candidates']
            single = next(p for p in c if self.m.products[p]['kind'] == J.KIND_SINGLE)
            pack = next(p for p in c if self.m.products[p]['kind'] == J.KIND_SET)
            e1 = self.store.register(self.sid, 'テスト', 'MAC-1', 'A-01-01', jan, 1, chosen_pid=single)
            e2 = self.store.register(self.sid, 'テスト', 'MAC-1', 'A-01-01', jan, 2, chosen_pid=pack)
            e3 = self.store.register(self.sid, 'テスト', 'MAC-1', 'A-01-01', jan, 3)
            self.assertEqual([e['数量単位'] for e in (e1, e2, e3)], [S.UNIT_SINGLE, S.UNIT_SET, S.UNIT_PIECE], jan)
            expect_single[single] = expect_single.get(single, 0) + 1
            expect_set[pack] = expect_set.get(pack, 0) + 2
        r = A.run([self.sid], self.store, set_composition={})
        self.assertEqual({b['pid']: b['qty'] for b in r['products']}, expect_single)
        self.assertEqual({b['pid']: b['qty'] for b in r['set_pending']}, expect_set)
        self.assertEqual({b['jan']: (b['qty'], b['reason']) for b in r['unregistered']},
                         {j: (3, A.CHK_SHARED_UNSELECTED) for j in self.shared})
        self.assertEqual(r['stats']['units'], len(self.shared))               # 単品に入るのはバラの1個だけ
        self.assertEqual(A.run([self.sid], self.store, set_composition={})['fingerprint'], r['fingerprint'])

    def test_crystal_h2o_rehearsal_case(self):
        r = self.m.resolve_jan('4544059033092')
        self.assertEqual((r['status'], r['candidates']), (J.R_SHARED, ['P000829', 'P000814']))
        self.store.register(self.sid, 'テスト', 'MAC-1', 'A-01-03', '4544059033092', 104, chosen_pid='P000829')
        self.store.register(self.sid, 'テスト', 'MAC-1', 'A-01-03', '4544059033092', 8, chosen_pid='P000814')
        a = A.run([self.sid], self.store, set_composition={})
        self.assertEqual([(b['pid'], b['qty']) for b in a['products']], [('P000829', 104)])
        self.assertEqual([(b['pid'], b['qty']) for b in a['set_pending']], [('P000814', 8)])


# ─────────────── G. HTTP ───────────────

class Http(Base):
    def start(self, host='127.0.0.1', token=''):
        app = APP.App(self.store, token)
        self.httpd = ThreadingHTTPServer((host, 0), APP.make_handler(app, host, 0))
        port = self.httpd.server_address[1]
        self.httpd.RequestHandlerClass = APP.make_handler(app, host, port)
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()
        self.base = f'http://127.0.0.1:{port}'
        self.addCleanup(self.httpd.shutdown)

    def call(self, path, body=None, headers=None, ctype='application/json'):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(self.base + path, data=data, headers=dict(headers or {}))
        if data is not None:
            req.add_header('Content-Type', ctype)
        try:
            with urllib.request.urlopen(req) as res:
                return res.status, json.loads(res.read() or b'{}') if 'json' in res.headers['Content-Type'] else {}
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read() or b'{}')

    def test_flow_scan_qty_undo(self):
        self.start()
        who = {'session': self.sid, 'staff': '英樹', 'device': 'MAC-1'}
        self.assertEqual(self.call('/')[0], 200)
        st, r = self.call('/api/lookup', {'code': GOAT})
        self.assertEqual((st, r['status'], r['product']['pid'], r['product']['cost']), (200, J.R_MATCH, 'P000001', 1270))
        st, r = self.call('/api/lookup', {'code': TOWEL})
        self.assertEqual([c['pid'] for c in r['candidates']], ['P000002', 'P000003'])
        st, r = self.call('/api/register', dict(who, raw=GOAT, qty='30', location='A-01-01'))
        self.assertEqual((st, r['event']['数量']), (200, '30'))
        st, r = self.call('/api/register', dict(who, raw=GOAT, qty=GOAT))
        self.assertEqual((st, r['code']), (400, 'qty_looks_like_code'))
        st, r = self.call('/api/cancel', who)
        self.assertEqual(r['event']['イベント種別'], S.EV_CANCEL)
        st, r = self.call('/api/aggregate', who)
        self.assertEqual((st, r['stats']['cancelled'], r['stats']['units']), (200, 1, 0))
        self.assertTrue(os.path.exists(r['file']))

    def test_shared_jan_over_http(self):
        self.start()
        who = {'session': self.sid, 'staff': '英樹', 'device': 'MAC-1'}
        st, r = self.call('/api/lookup', {'code': SHARED})
        self.assertEqual((r['status'], r['product']), (J.R_SHARED, None))
        self.assertEqual([(c['pid'], c['kind']) for c in r['candidates']], [('P000005', '単品'), ('P000006', 'セット')])
        st, r = self.call('/api/register', dict(who, raw=SHARED, qty='8', chosen_pid='P000006'))
        self.assertEqual((st, r['event']['数量単位'], r['event']['内部管理ID']), (200, S.UNIT_SET, 'P000006'))

    def test_host_header_guard_on_localhost(self):
        self.start()
        st, _ = self.call('/api/info', headers={'Host': 'evil.example:80'})
        self.assertEqual(st, 403)

    def test_token_required_in_lan_mode(self):
        self.start(host='0.0.0.0', token='secret-token')
        self.assertEqual(self.call('/api/info')[0], 401)
        self.assertEqual(self.call('/api/info', headers={'X-Stocktake-Token': 'secret-token'})[0], 200)

    def test_json_only(self):
        self.start()
        self.assertEqual(self.call('/api/lookup', {'code': GOAT}, ctype='text/plain')[0], 415)


if __name__ == '__main__':
    unittest.main()
