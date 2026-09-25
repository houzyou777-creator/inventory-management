# -*- coding: utf-8 -*-
"""stocktake_aggregate.py — 棚卸しスキャン記録の集計(記録ファイル + 商品マスター → 集計xlsx)

使い方:
    python3 stocktake_aggregate.py <セッションID> [<セッションID> ...] [--data-dir DIR] [--out-dir DIR]

**読み取り専用。** 記録ファイル・商品マスターは読むだけ。書き出すのは新規の集計 xlsx 1本だけ
(既定: SourceData/Stocktake/summaries/)。在庫管理テーブル・Amazon・楽天・プライスターへは何も反映しない。

集計の規則:
    ・取消イベントの対象になった登録は数えない(行は明細に残す)
    ・JAN は集計のたびに商品マスターで照合し直す(記録時に未登録でも、後でマスターに登録すれば解決される)
      候補選択した登録は、選んだPが今も候補に含まれる限りその選択を使う
    ・実棚数量は物理的な単品P単位。セットPの登録は「セット換算待ち」に分け、単品へ混ぜない
      (セット構成表ができたら set_composition を渡すだけで 単品数 += セット数 × 入数 として合算される)
    ・同じ入力(記録・マスター・紐付け・構成表)からは必ず同じ結果になる。集計結果ハッシュで確かめられる
"""
import argparse
import hashlib
import json
import os
import sys
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import jan_lookup as J  # noqa: E402
import stocktake_store as S  # noqa: E402

NO_LOC = '(未指定)'

CHK_DUP_UNSELECTED = 'JAN重複(候補未選択)'
CHK_SET_UNSELECTED = 'セット品JAN(候補未選択)'
CHK_SHARED_UNSELECTED = '単品セット共通JAN(バラ/セット未選択)'
CHK_CHANGED = '照合結果が記録時と変化'
CHK_RESOLVED_LATER = '記録後に照合できた'
CHK_LARGE = '大量数量(確認済みで登録)'
CHK_REPEAT = '同じ場所・同じ数量の複数登録(二重カウントの可能性)'
CHK_BAD_CANCEL = '取消対象が見つからない'
CHK_BAD_ROW = 'データ異常'
CHK_BAD_LINK = '仮IDの紐付け先がマスターに無い'
CHK_NO_COST = '標準原価なし(金額を出せない)'


def load_set_composition():
    """セット構成表(セットP → [(単品P, 入数)])。

    セット構成表は別レーンで設計中(2026-09-25)。Step 1 では空を返し、セットPの登録は
    「セット換算待ち」に残す。構成表ができたらここで読み込むだけで換算が有効になる。
    """
    return {}


def _loc_text(counter):
    return ' / '.join(f'{k}:{v}' for k, v in sorted(counter.items()))


def aggregate(events, master, temp_links, temp_items, config, set_composition=None):
    comp = set_composition or {}
    regs = [e for e in events if e['イベント種別'] == S.EV_REGISTER]
    reg_ids = {e['イベントID'] for e in regs}
    cancelled, checks, details = set(), [], []

    for e in events:
        if e['イベント種別'] == S.EV_CANCEL:
            t = e['取消対象イベントID']
            if t in reg_ids:
                cancelled.add(t)
            else:
                checks.append([CHK_BAD_CANCEL, e['イベントID'], f'取消対象 {t!r} が記録にありません', '記録ファイルを確認'])
            details.append((e, '取消イベント', ''))

    products = {}
    unregistered, temps, set_pending = {}, {}, {}
    seen = defaultdict(list)   # (P/JAN, ロケーション, 数量) → イベントID(二重カウントの疑い)

    def bucket(store, key, init):
        if key not in store:
            store[key] = dict(init, qty=0, count=0, locs=Counter(), events=[], staff=set(), first='', last='')
        return store[key]

    def add(b, e, q, loc):
        b['qty'] += q
        b['count'] += 1
        b['locs'][loc] += q
        b['events'].append(e['イベントID'])
        b['staff'].add(e['担当者'])
        b['first'] = min(b['first'] or e['記録日時'], e['記録日時'])
        b['last'] = max(b['last'], e['記録日時'])

    def add_single(pid, e, q, loc, via_set=''):
        p = master.products[pid]
        b = bucket(products, pid, {'pid': pid, 'jan': str(p['jan_raw'] or ''), 'name': p['name'],
                                   'cost': p['cost'], 'direct': 0, 'from_sets': 0, 'set_detail': Counter()})
        add(b, e, q, loc)
        if via_set:
            b['from_sets'] += q
            b['set_detail'][via_set] += q
        else:
            b['direct'] += q

    for e in regs:
        eid = e['イベントID']
        if eid in cancelled:
            details.append((e, '取消済(数えない)', ''))
            continue
        try:
            q = int(e['数量'])
        except ValueError:
            checks.append([CHK_BAD_ROW, eid, f'数量 {e["数量"]!r} が整数ではない', '記録ファイルを確認'])
            details.append((e, 'データ異常(数えない)', ''))
            continue
        loc = e['ロケーション'] or NO_LOC
        code, kind, snap = e['正規化コード'], e['コード種別'], e['内部管理ID']
        pid, how = '', ''
        if kind == J.CODE_JAN:
            res = master.resolve_jan(code)
            if res['status'] == J.R_MATCH:
                pid = res['pid']
            elif res['status'] in (J.R_DUPLICATE, J.R_SET_ONLY, J.R_SHARED):
                if e['入力方法'] == S.METHOD_CHOSEN and snap in res['candidates']:
                    pid = snap
                elif res['status'] == J.R_SET_ONLY and res['pid']:
                    pid = res['pid']
                else:
                    # 選ばれていない登録は実棚に入れない。共通JANを単品と決めつけると過少計上になる
                    label = {J.R_DUPLICATE: CHK_DUP_UNSELECTED, J.R_SET_ONLY: CHK_SET_UNSELECTED,
                             J.R_SHARED: CHK_SHARED_UNSELECTED}[res['status']]
                    checks.append([label, eid, f'JAN {code} 数量 {q} 候補: {", ".join(res["candidates"])}',
                                   '現物を確認し、候補から商品を特定する(明細の数量は実棚に未算入)'])
                    if snap:
                        # 修正前に単品へ自動確定された記録など。記録時の判断を黙って使わない
                        checks.append([CHK_CHANGED, eid, f'記録時 {snap} → 現在は {res["status"]}(要選択)',
                                       '現物がバラか梱包済みセットかを確認'])
                    b = bucket(unregistered, code, {'jan': code, 'reason': label})
                    add(b, e, q, loc)
                    details.append((e, '要確認(実棚に未算入)', ''))
                    continue
            else:
                b = bucket(unregistered, code, {'jan': code, 'reason': J.R_UNREGISTERED})
                add(b, e, q, loc)
                details.append((e, '未登録JAN(実棚に未算入)', ''))
                if snap:
                    checks.append([CHK_CHANGED, eid, f'記録時 {snap} → 現在は未登録', '商品マスターのJAN変更を確認'])
                continue
        elif kind == J.CODE_TEMP:
            link = temp_links.get(code, '')
            if link and link in master.products:
                pid, how = link, f'仮ID {code} の紐付け'
            else:
                if link:
                    checks.append([CHK_BAD_LINK, eid, f'{code} → {link}', 'temp_links.csv を確認'])
                memo = temp_items.get(code, {}).get('メモ', '')
                b = bucket(temps, code, {'temp': code, 'memo': memo})
                add(b, e, q, loc)
                details.append((e, '仮ID(実棚に未算入)', ''))
                continue
        else:
            checks.append([CHK_BAD_ROW, eid, f'コード種別 {kind!r} は集計できない', '記録ファイルを確認'])
            details.append((e, 'データ異常(数えない)', ''))
            continue

        if snap and snap != pid:
            checks.append([CHK_CHANGED, eid, f'記録時 {snap} → 現在 {pid}', '商品マスターの変更内容を確認'])
        elif not snap and pid and kind == J.CODE_JAN:
            checks.append([CHK_RESOLVED_LATER, eid, f'JAN {code} → {pid}', '対応不要(記録後にマスター登録された)'])
        if q >= config['quantity']['confirm_threshold']:
            checks.append([CHK_LARGE, eid, f'{pid} 数量 {q}', '現物の数量を再確認'])

        p = master.products[pid]
        if p['kind'] == J.KIND_SINGLE:
            add_single(pid, e, q, loc)
            seen[(pid, loc, q)].append(eid)
            details.append((e, '実棚に算入', pid + (f'({how})' if how else '')))
        elif pid in comp:
            for spid, per in comp[pid]:
                add_single(spid, e, q * per, loc, via_set=f'{pid}×{q}(入数{per})')
            details.append((e, 'セット換算して実棚に算入', pid))
        else:
            b = bucket(set_pending, pid, {'pid': pid, 'jan': str(p['jan_raw'] or ''), 'name': p['name']})
            add(b, e, q, loc)
            details.append((e, 'セット換算待ち(実棚に未算入)', pid))

    for (key, loc, q), ids in sorted(seen.items()):
        if len(ids) > 1:
            checks.append([CHK_REPEAT, ', '.join(ids), f'{key} {loc} 数量{q} × {len(ids)}回',
                           '別の現物なら対応不要。同じ現物を数えていたら1件を取消'])
    for pid, b in sorted(products.items()):
        if b['cost'] is None:
            checks.append([CHK_NO_COST, '', f'{pid} {b["name"][:30]}', '商品マスターの標準原価を確認(実棚数量には影響しない)'])

    active = [e for e in regs if e['イベントID'] not in cancelled]
    result = {
        'products': [products[k] for k in sorted(products)],
        'unregistered': [unregistered[k] for k in sorted(unregistered)],
        'temps': [temps[k] for k in sorted(temps)],
        'set_pending': [set_pending[k] for k in sorted(set_pending)],
        'checks': checks,
        'details': sorted(details, key=lambda d: d[0]['イベントID']),
        'stats': {'events': len(events), 'registrations': len(regs), 'cancelled': len(cancelled),
                  'active': len(active), 'products': len(products),
                  'units': sum(b['qty'] for b in products.values()),
                  'unregistered_qty': sum(b['qty'] for b in unregistered.values()),
                  'temp_qty': sum(b['qty'] for b in temps.values()),
                  'set_pending_qty': sum(b['qty'] for b in set_pending.values())},
    }
    result['fingerprint'] = fingerprint(result)
    return result


def fingerprint(result):
    """集計結果の指紋。同じ入力から同じ集計になることを確かめるために使う。"""
    def conv(o):
        if isinstance(o, (set, frozenset)):
            return sorted(o)
        if isinstance(o, Counter):
            return sorted(o.items())
        if isinstance(o, tuple):
            return list(o)
        return str(o)
    body = {k: v for k, v in result.items() if k != 'fingerprint'}
    body['details'] = [(d[0]['イベントID'], d[1], d[2]) for d in result['details']]
    return hashlib.sha256(json.dumps(body, ensure_ascii=False, sort_keys=True, default=conv).encode()).hexdigest()


def run(session_ids, store, set_composition=None):
    events = []
    for sid in session_ids:
        if not os.path.exists(store.events_path(sid)):
            raise S.StoreError(f'セッション {sid} の記録がありません', 'no_session')
        events += store.events(sid)
    comp = load_set_composition() if set_composition is None else set_composition
    return aggregate(events, store.master, store.temp_links(), store.temp_items(), store.config, comp)


# ─────────────────────────────── 出力 ───────────────────────────────

def _sheet(wb, title, header, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDEBF7')
    for r in rows:
        ws.append(r)
    ws.freeze_panes = 'A2'
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    return ws


def write_xlsx(result, session_ids, store, out_dir):
    wb = Workbook()
    rows = [[b['pid'], b['jan'], b['name'], b['qty'], b['direct'], b['from_sets'],
             _loc_text(b['set_detail']), b['count'], _loc_text(b['locs']),
             b['cost'] if b['cost'] is not None else '',
             b['qty'] * b['cost'] if b['cost'] is not None else '', b['last'], ', '.join(sorted(b['staff']))]
            for b in result['products']]
    _sheet(wb, '実棚集計', ['内部管理ID', 'JAN', '商品名', '実棚数量(単品)', 'うち単品で計数', 'うちセット換算',
                           'セット換算内訳', '登録回数', 'ロケーション内訳', '標準原価', '金額(参考)', '最終登録', '担当者'],
           rows, [11, 15, 50, 12, 11, 11, 20, 8, 28, 9, 11, 19, 14], first=True)

    def simple(buckets, lead):
        return [lead(b) + [b['qty'], b['count'], _loc_text(b['locs']), b['first'], b['last'],
                           ', '.join(sorted(b['staff'])), ', '.join(b['events'])] for b in buckets]
    tail = ['数量', '登録回数', 'ロケーション内訳', '初回登録', '最終登録', '担当者', 'イベントID']
    tw = [8, 8, 28, 19, 19, 14, 40]
    _sheet(wb, '未登録JAN', ['JAN', '理由'] + tail, simple(result['unregistered'], lambda b: [b['jan'], b['reason']]),
           [15, 22] + tw)
    _sheet(wb, '仮ID', ['仮ID', 'メモ'] + tail, simple(result['temps'], lambda b: [b['temp'], b['memo']]), [11, 36] + tw)
    _sheet(wb, 'セット換算待ち', ['内部管理ID(セット)', 'JAN', '商品名', 'セット数'] + tail[1:],
           simple(result['set_pending'], lambda b: [b['pid'], b['jan'], b['name']]), [16, 15, 50] + tw)
    _sheet(wb, '要確認', ['種別', 'イベントID', '内容', '推奨対応'], result['checks'], [34, 26, 60, 50])
    det = [[e[c] for c in S.EVENT_COLUMNS] + [how, pid] for e, how, pid in result['details']]
    _sheet(wb, '明細', S.EVENT_COLUMNS + ['集計上の扱い', '集計先'], det, [26, 16] + [12] * 18 + [24, 16])

    st = result['stats']
    info = [
        ['セッションID', ', '.join(session_ids)],
        ['集計日時', datetime.now().strftime(S.TIME_FMT)],
        ['商品マスター', store.master.path],
        ['商品マスター sha256', store.master.sha256],
        ['記録ファイル sha256', ' / '.join(f'{sid}: {J.sha256_file(store.events_path(sid))}' for sid in session_ids)],
        ['集計結果ハッシュ', result['fingerprint']],
        ['イベント数(登録/取消を含む)', st['events']],
        ['登録', st['registrations']],
        ['うち取消済', st['cancelled']],
        ['集計対象の登録', st['active']],
        ['実棚: 商品数(単品P)', st['products']],
        ['実棚: 数量合計(単品)', st['units']],
        ['未登録JAN・候補未選択の数量', st['unregistered_qty']],
        ['仮IDの数量', st['temp_qty']],
        ['セット換算待ちのセット数', st['set_pending_qty']],
        ['要確認', len(result['checks'])],
        ['注記', '実棚数量は物理的な単品P単位。帳簿在庫・在庫管理テーブル・Amazon・楽天・プライスターへは反映していない'],
        ['注記', 'このファイルは記録ファイルと商品マスターから再生成できる(同じ入力なら集計結果ハッシュが一致する)'],
    ]
    _sheet(wb, '集計情報', ['項目', '値'], info, [30, 110])
    os.makedirs(out_dir, exist_ok=True)
    name = '_'.join(session_ids) if len(session_ids) <= 2 else f'{session_ids[0]}_plus{len(session_ids) - 1}'
    out = os.path.join(out_dir, f'stocktake_summary_{name}_{datetime.now():%Y%m%d_%H%M%S}.xlsx')
    wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description='棚卸しスキャン記録の集計(読み取り専用・集計xlsxを新規作成)')
    ap.add_argument('sessions', nargs='+')
    ap.add_argument('--data-dir', default=S.DATA_DIR)
    ap.add_argument('--out-dir', default='')
    a = ap.parse_args()
    store = S.Store(J.load_master(), J.load_config(), data_dir=a.data_dir, writer=False)
    result = run(a.sessions, store)
    out = write_xlsx(result, a.sessions, store, a.out_dir or os.path.join(a.data_dir, S.SUMMARY_DIR))
    st = result['stats']
    print(f'実棚 {st["products"]}商品 / {st["units"]}個(単品)  未登録・未選択 {st["unregistered_qty"]}  '
          f'仮ID {st["temp_qty"]}  セット換算待ち {st["set_pending_qty"]}  要確認 {len(result["checks"])}')
    print('集計結果ハッシュ:', result['fingerprint'])
    print('出力:', out)


if __name__ == '__main__':
    main()
