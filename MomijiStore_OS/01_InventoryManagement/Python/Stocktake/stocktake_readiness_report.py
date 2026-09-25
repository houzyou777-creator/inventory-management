# -*- coding: utf-8 -*-
"""stocktake_readiness_report.py — JANスキャン棚卸し Step 0: P番号整合 + JAN整備レポート

使い方:
    python3 stocktake_readiness_report.py [--out-dir <出力フォルダ>]

**読み取り専用。** 商品マスター・在庫管理テーブル・旧システムは read_only で開き、保存しない。
書き出すのは新規のレポート xlsx 1本だけ(既定: SourceData/Output/)。

なぜ必要か(2026-09-25 承認):
    ・P番号の意味が一意でないまま照合すると、別商品へ数量が入る。最優先で確かめる
      (旧 Excel/在庫管理システム_v1.0.xlsm は P000001=ロイヤルカナン、正の商品マスターは P000001=クリニーク)
    ・スキャンで特定できない単品(JAN重複・空欄・ASIN混入)を棚卸し前に把握する
"""
import argparse
import os
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import jan_lookup as J  # noqa: E402
from jan_lookup import MP  # noqa: E402

STOCK_FILE = os.path.join(MP.INV_SD, '在庫管理テーブル_v1.1.xlsm')
SH_STOCK = '在庫管理テーブル'
STOCK_HEAD = {'sid': '在庫ID', 'pid': '内部管理ID', 'jan': 'JAN', 'name': '商品名', 'zone': '保管区分'}
SH_LISTING = '出品テーブル'
LISTING_HEAD = {'lid': '出品ID', 'pid': '内部管理ID', 'asin': 'ASIN'}
SH_LEGACY = '商品マスター'

ST_SCANNABLE = 'スキャン可'
ST_DUP = 'JAN重複(候補選択が必要)'
ST_SHARED = '単品セット共通JAN(バラ/セットを選択)'
ST_BLANK = 'JAN空欄(仮ID運用)'
ST_ASIN = 'JAN欄にASIN(要修正)'
ST_BAD = 'JAN形式エラー(要修正)'

PASS, NG, WARN, INFO = 'PASS', 'NG', '要確認', '参考'   # NG だけが「照合に使えない」を意味する
FILL = {PASS: 'C6EFCE', NG: 'FFC7CE', WARN: 'FFEB9C', INFO: 'DDEBF7'}


def norm_name(s):
    return ''.join(unicodedata.normalize('NFKC', str(s or '')).split())


def _rows(path, sheet, wanted, match_contains=False):
    """見出し名で列を引いて dict のリストを返す(read_only・保存しない)。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        header = [str(c or '').strip() for c in next(it)]
        pos = {}
        for k, h in wanted.items():
            idx = [i for i, x in enumerate(header) if (h in x if match_contains else h == x)]
            if not idx:
                raise J.MasterError(f'{os.path.basename(path)}「{sheet}」に見出し「{h}」がありません')
            pos[k] = idx[0]
        out = []
        for i, r in enumerate(it, start=2):
            rec = {k: r[p] if p < len(r) else None for k, p in pos.items()}
            if all(v in (None, '') for v in rec.values()):
                continue
            rec['_row'] = i
            out.append(rec)
        return out
    finally:
        wb.close()


def s(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return '' if v is None else str(v).strip()


# ─────────────────────────────── P番号整合 ───────────────────────────────

def check_ids(master, stock, listing, legacy):
    checks, legacy_rows = [], []
    P = master.products

    def add(name, ok, detail, level=None):
        checks.append((name, level or (PASS if ok else NG), detail))

    add('商品マスター: 内部管理IDの形式(P+6桁)と重複', True,
        f'{len(P)}件すべて形式どおり・重複0(load_master が検証。1件でも違反があれば停止する)')

    missing = [r for r in stock if s(r['pid']) not in P]
    add('在庫管理テーブル: 内部管理IDがすべて商品マスターに存在', not missing,
        f'{len(stock)}行中 不在 {len(missing)}' + (f' 例: {[s(r["pid"]) for r in missing[:5]]}' if missing else ''))
    jan_diff = []
    for r in stock:
        p = P.get(s(r['pid']))
        if p and (J.master_jan_key(r['jan'])[0] or s(r['jan'])) != (p['jan'] or s(p['jan_raw'])):
            jan_diff.append((s(r['pid']), s(r['jan']), s(p['jan_raw'])))
    add('在庫管理テーブル: JANが商品マスターと一致', not jan_diff,
        f'不一致 {len(jan_diff)}' + (f' 例: {jan_diff[:3]}' if jan_diff else ''))
    name_diff = [f'{s(r["pid"])} 在庫表「{s(r["name"])}」/ マスター「{P[s(r["pid"])]["name"]}」'
                 for r in stock if s(r['pid']) in P and norm_name(r['name']) != norm_name(P[s(r['pid'])]['name'])]
    # 照合は商品マスターだけで行うので表示名の差は照合に影響しない。ただし在庫表への反映(Phase 4)前に直す
    add('在庫管理テーブル: 商品名が商品マスターと一致(空白・全半角の差は無視)', True,
        f'不一致 {len(name_diff)}' + (' ' + ' ／ '.join(name_diff[:5]) if name_diff else ''),
        level=WARN if name_diff else PASS)
    sid_dup = [k for k, v in Counter(s(r['sid']) for r in stock).items() if v > 1]
    key_dup = [k for k, v in Counter((s(r['pid']), s(r['zone'])) for r in stock).items() if v > 1]
    add('在庫管理テーブル: 在庫ID・(内部管理ID×保管区分) の重複なし', not sid_dup and not key_dup,
        f'在庫ID重複 {len(sid_dup)} / 内部管理ID×保管区分 重複 {len(key_dup)}')
    set_rows = [s(r['pid']) for r in stock if s(r['pid']) in P and P[s(r['pid'])]['kind'] != J.KIND_SINGLE]
    add('在庫管理テーブル: 在庫行は単品Pだけ(物理単品基準と整合)', not set_rows,
        f'単品以外の在庫行 {len(set_rows)}' + (f' 例: {set_rows[:5]}' if set_rows else ''))

    lmiss = [r for r in listing if s(r['pid']) not in P]
    add('出品テーブル: 内部管理IDがすべて商品マスターに存在', not lmiss,
        f'{len(listing)}行中 不在 {len(lmiss)}' + (f' 例: {[s(r["lid"]) for r in lmiss[:5]]}' if lmiss else ''))
    ldup = [k for k, v in Counter(s(r['lid']) for r in listing).items() if v > 1]
    add('出品テーブル: 出品IDの重複なし', not ldup, f'重複 {len(ldup)}')

    # 旧システム: 同じP番号が別の商品を指しているか
    by_code = defaultdict(list)       # マスター側の JAN欄の値(ASIN混入も含む)・出品ASIN → pid
    for pid, p in P.items():
        if s(p['jan_raw']):
            by_code[s(p['jan_raw']).upper()].append(pid)
    for r in listing:
        if s(r['asin']):
            by_code[s(r['asin']).upper()].append(s(r['pid']))
    conflict = 0
    for r in legacy:
        pid = s(r['pid'])
        mp = P.get(pid)
        same = bool(mp) and norm_name(mp['name']) == norm_name(r['name'])
        conflict += 0 if same else 1
        hits = sorted({x for c in (s(r['jan']).upper(), s(r['asin']).upper()) if c for x in by_code.get(c, [])})
        legacy_rows.append([pid, s(r['name']), s(r['jan']), s(r['asin']),
                            mp['name'] if mp else '(商品マスターに無し)',
                            '同一' if same else '別商品', ', '.join(hits) or '(該当なし)'])
    add('旧 在庫管理システム_v1.0: P番号の意味が商品マスターと異なる', True,
        f'旧システム {len(legacy)}件中 {conflict}件で同じP番号が別商品を指す → 旧システムのIDは照合に使用しない',
        level=INFO)
    try:
        J.load_master(J.LEGACY_SYSTEM_FILE)
        add('照合コード: 旧システムを商品マスターとして読ませると停止する', False, '停止しなかった')
    except J.MasterError as e:
        add('照合コード: 旧システムを商品マスターとして読ませると停止する', True, str(e))
    return checks, legacy_rows


# ─────────────────────────────── JAN整備 ───────────────────────────────

def jan_status(master, stock_pids):
    issues = {pid: why for pid, _, why in master.jan_issues}
    rows = []
    for pid, p in master.products.items():
        if p['kind'] != J.KIND_SINGLE:
            continue
        if pid in issues:
            why = issues[pid]
            st = ST_BLANK if why == '空欄' else ST_ASIN if why == 'ASINが入っている' else ST_BAD
            note = why
        elif len(master.by_jan_single.get(p['jan'], [])) > 1:
            st, note = ST_DUP, '同JANの単品: ' + ', '.join(master.by_jan_single[p['jan']])
        elif p['jan'] in master.by_jan_set:
            st, note = ST_SHARED, '同JANのセット: ' + ', '.join(master.by_jan_set[p['jan']])
        else:
            st, note = ST_SCANNABLE, ''
        rows.append([pid, s(p['jan_raw']), p['jan'], st, '有' if pid in stock_pids else '無',
                     p['cost'] if p['cost'] is not None else '', p['name'], note])
    return rows


def set_only_rows(master):
    out = []
    for jan, pids in master.by_jan_set.items():
        if jan not in master.by_jan_single:
            for pid in pids:
                out.append([jan, pid, master.products[pid]['name'],
                            'セットのJANのみ一致。Step 1 では「セット品JAN」として記録し、単品換算はセット構成表の完成後'])
    return sorted(out)


# ─────────────────────────────── 出力 ───────────────────────────────

def write_sheet(wb, title, header, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDEBF7')
    for r in rows:
        ws.append(r)
    ws.freeze_panes = 'A2'
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    return ws


def build(out_dir):
    master = J.load_master()
    stock = _rows(STOCK_FILE, SH_STOCK, STOCK_HEAD)
    listing = _rows(J.MASTER_FILE, SH_LISTING, LISTING_HEAD)
    legacy = _rows(J.LEGACY_SYSTEM_FILE, SH_LEGACY,
                   {'pid': '内部管理ID', 'name': '商品名', 'jan': 'JANコード', 'asin': 'ASIN'}, match_contains=True)
    legacy = [r for r in legacy if s(r['pid'])]

    checks, legacy_rows = check_ids(master, stock, listing, legacy)
    stock_pids = {s(r['pid']) for r in stock}
    status = jan_status(master, stock_pids)
    cnt = Counter(r[3] for r in status)
    dup_groups = {k: v for k, v in master.by_jan_single.items() if len(v) > 1}
    shared = {k for k in master.by_jan_single if k in master.by_jan_set}
    sets_only = set_only_rows(master)
    ok_ids = all(c[1] != NG for c in checks)
    warns = sum(1 for c in checks if c[1] == WARN)

    singles = len(status)
    summary = [
        ['P番号の整合(最優先)', (PASS if ok_ids else NG) + (f'(要確認{warns})' if warns else ''),
         'シート「P番号整合」参照。NG があれば MVP の照合を使わない。要確認は照合に影響しない差'],
        ['商品マスター 件数', len(master.products), f'単品 {singles} / セット {len(master.products) - singles}'],
        ['商品マスター 版(sha256 先頭12桁)', master.version, master.path],
        ['単品: スキャンで1商品に特定できる', cnt[ST_SCANNABLE], f'{cnt[ST_SCANNABLE] / singles:.1%}'],
        ['単品: JAN重複(スキャン時に候補選択)', cnt[ST_DUP], f'{len(dup_groups)}グループ'],
        ['単品: JAN空欄(仮ID運用・後で紐付け)', cnt[ST_BLANK], ''],
        ['単品: JAN欄にASIN(要修正)', cnt[ST_ASIN], ''],
        ['単品: JAN形式エラー(要修正)', cnt[ST_BAD], ''],
        ['単品: 単品とセットで共通のJAN(バラ/セットを選択)', cnt[ST_SHARED],
         f'{len(shared)}JAN。単品へ自動確定しない。梱包済みセットはセットP×セット数で「セット換算待ち」'],
        ['セットだけに一致するJAN', len({r[0] for r in sets_only}), '「セット品JAN」として記録。換算はセット構成表の完成後'],
        ['単品で在庫管理テーブルに行が無い', sum(1 for r in status if r[4] == '無'),
         'スキャン・集計は可能。在庫管理テーブルへの反映(Phase 4)前に在庫行の追加要否を決める'],
        ['在庫管理テーブル 行数', len(stock), ''],
        ['注記', '', '読み取り専用で作成。正本(商品マスター・在庫管理テーブル)には書き込んでいない'],
    ]

    wb = Workbook()
    write_sheet(wb, '概要', ['項目', '値', '備考'], summary, [40, 16, 90], first=True)
    ws = write_sheet(wb, 'P番号整合', ['チェック', '結果', '詳細'], checks, [52, 8, 110])
    for r in ws.iter_rows(min_row=2):
        r[1].fill = PatternFill('solid', fgColor=FILL[r[1].value])
        r[2].alignment = Alignment(wrap_text=True)
    write_sheet(wb, '旧システムID比較', ['旧システムP番号', '旧システム商品名', '旧JAN', '旧ASIN', '商品マスター同P番号の商品名',
                                       '判定', '商品マスターでの該当P(JAN/ASINで照合)'], legacy_rows, [14, 40, 16, 14, 50, 8, 26])
    order = [ST_DUP, ST_SHARED, ST_ASIN, ST_BAD, ST_BLANK, ST_SCANNABLE]
    status.sort(key=lambda r: (order.index(r[3]), r[0]))
    head = ['内部管理ID', 'JAN(マスター値)', '正規化JAN', 'スキャン状態', '在庫行', '標準原価', '商品名', '備考']
    widths = [12, 16, 15, 26, 8, 10, 60, 60]
    write_sheet(wb, '単品JAN状態', head, status, widths)
    dup_rows = [[jan, pid, s(master.products[pid]['jan_raw']), master.products[pid]['name']]
                for jan, pids in sorted(dup_groups.items()) for pid in pids]
    write_sheet(wb, 'JAN重複_単品', ['正規化JAN', '内部管理ID', 'JAN(マスター値)', '商品名'], dup_rows, [15, 12, 16, 80])
    write_sheet(wb, 'セットのみ一致JAN', ['正規化JAN', '内部管理ID(セット)', '商品名', '扱い'], sets_only, [15, 16, 70, 70])
    write_sheet(wb, '在庫行なし_単品', head, [r for r in status if r[4] == '無'], widths)

    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, f'stocktake_jan_readiness_{datetime.now():%Y%m%d_%H%M}.xlsx')
    wb.save(out)
    return out, summary, checks


def main():
    ap = argparse.ArgumentParser(description='JANスキャン棚卸し Step 0: P番号整合 + JAN整備レポート(読み取り専用)')
    ap.add_argument('--out-dir', default=MP.INV_OUTPUT)
    a = ap.parse_args()
    out, summary, checks = build(a.out_dir)
    print('== P番号整合 ==')
    for name, res, detail in checks:
        print(f'  [{res}] {name}\n         {detail}')
    print('== 概要 ==')
    for k, v, note in summary:
        print(f'  {k}: {v}  {note}')
    print('出力:', out)


if __name__ == '__main__':
    main()
