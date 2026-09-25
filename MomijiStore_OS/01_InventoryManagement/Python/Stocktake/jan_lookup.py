# -*- coding: utf-8 -*-
"""jan_lookup.py — JANスキャン棚卸しの共通部品: コードの正規化と商品マスター索引

**読み取り専用。** 商品マスターは read_only で開き、保存しない。

なぜ独立した部品にするか:
    スキャン画面(記録時)と集計(後から再照合)が「同じ規則」で JAN → 内部管理ID を解決しないと、
    生データから集計を再現できなくなる。規則をここ1か所に置く。

照合の原則(Master_Design §2):
    ・主キーは内部管理ID(P連番)。JAN は「照合キー」であって一意ではない
      (例: JAN 4533141407514 に 50本/100本/200本入りの3商品)。重複時は推測せず候補を返す
    ・棚卸の基準は物理的な単品P。JAN が単品にもセットにも一致したら単品を採る
    ・正は `SourceData/商品マスター_単品_v1.0.xlsx` のみ。
      旧 `Excel/在庫管理システム_v1.0.xlsm` は別のID体系(サンプルデータ)なので照合に使わせない
"""
import hashlib
import json
import os
import re
import sys
import unicodedata
from collections import defaultdict

from openpyxl import load_workbook

# 共通モジュール(02_Analytics/Python/momiji_paths.py)の場所は __file__ からの相対位置で求める
sys.path.insert(0, os.path.normpath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', '02_Analytics', 'Python')))
import momiji_paths as MP  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(HERE, 'stocktake_config.json')

MASTER_FILE = os.path.join(MP.INV_SD, '商品マスター_単品_v1.0.xlsx')
SH_MASTER = '商品マスター'
# 列は位置ではなく見出し名で引く(列が足されても壊れないように)
MASTER_HEAD = {'pid': '内部管理ID', 'jan': 'JAN', 'name': '商品名', 'kind': '商品種別', 'cost': '標準原価'}

# 旧システム(ID体系が違う)。誤って指定されたら止める
LEGACY_SYSTEM_FILE = os.path.join(MP.OS_ROOT_STR, '01_InventoryManagement', 'Excel', '在庫管理システム_v1.0.xlsm')

PID_RE = re.compile(r'^P\d{6}$')
KIND_SINGLE = '単品'
KIND_SET = 'セット'

# コード種別(スキャンした文字列が何か)
CODE_JAN = 'JAN'
CODE_CASE = 'ケースコード'
CODE_INVALID = '形式エラー'
CODE_TEMP = '仮ID'
CODE_LOCATION = 'ロケーション'
CODE_COMMAND = '操作コード'

# 照合結果
R_MATCH = '一致'
R_CHOSEN = '候補選択'
R_DUPLICATE = 'JAN重複'
R_SET_ONLY = 'セット品JAN'
R_UNREGISTERED = '未登録JAN'
R_TEMP = '仮ID'


class MasterError(RuntimeError):
    """商品マスターを安全に使えない(処理を続けさせない)"""


def load_config(path=CONFIG_FILE):
    with open(path, encoding='utf-8') as f:
        return json.load(f)


# ─────────────────────────────── コードの正規化 ───────────────────────────────

def gtin_check_ok(digits):
    """GTIN-8/12/13/14 共通のチェックディジット検証(右から 3,1,3,1… の重み)。"""
    body, check = digits[:-1], int(digits[-1])
    total = sum(int(d) * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(body)))
    return (10 - total % 10) % 10 == check


def _clean(raw):
    # iPad の日本語入力で全角数字が入っても同じ値になるよう NFKC で揃える
    s = unicodedata.normalize('NFKC', str(raw or '')).strip()
    return re.sub(r'\s+', '', s).upper()


def normalize_code(raw, config):
    """スキャン/手入力された文字列を分類し、照合に使う正規化コードを返す。

    戻り値: {'kind', 'code', 'message'}
      JAN は 13桁(EAN-13) または 8桁(EAN-8)に揃える。UPC-A(12桁)は先頭0を補って13桁にする
      (マスターの 192333006122 と楽天の 0192333006122 を同じ商品として扱うため)。
    """
    s = _clean(raw)
    if not s:
        return {'kind': CODE_INVALID, 'code': '', 'message': '空の入力です'}
    if s in config['commands']:
        return {'kind': CODE_COMMAND, 'code': s, 'message': config['commands'][s]}
    prefix = config['temp_id_prefix']
    if re.fullmatch(re.escape(prefix) + r'\d{%d}' % config['temp_id_digits'], s):
        return {'kind': CODE_TEMP, 'code': s, 'message': ''}
    if re.fullmatch(config['location_pattern'], s):
        return {'kind': CODE_LOCATION, 'code': s, 'message': ''}
    if not s.isdigit():
        return {'kind': CODE_INVALID, 'code': s, 'message': 'JAN・ロケーション・仮IDのどれでもありません'}
    n = len(s)
    if n in (8, 12, 13, 14) and not gtin_check_ok(s):
        return {'kind': CODE_INVALID, 'code': s,
                'message': 'チェックディジットが合いません(読み取りミス・打ち間違いの可能性)'}
    if n == 13 or n == 8:
        return {'kind': CODE_JAN, 'code': s, 'message': ''}
    if n == 12:
        return {'kind': CODE_JAN, 'code': '0' + s, 'message': 'UPC-A を13桁へ正規化'}
    if n == 14 and s[0] == '0':
        # 梱包インジケータ0の GTIN-14 は単品と同じ商品を指す
        return {'kind': CODE_JAN, 'code': s[1:], 'message': 'GTIN-14(インジケータ0)を13桁へ正規化'}
    if n in (14, 16):
        # 外箱(ケース)のコード。1箱を1個として数える事故を防ぐため登録させない
        return {'kind': CODE_CASE, 'code': s,
                'message': '外箱(ケース)のコードです。中の単品のJANを読んでください'}
    return {'kind': CODE_INVALID, 'code': s, 'message': f'{n}桁のコードはJANとして扱えません'}


def master_jan_key(value):
    """商品マスターのJAN欄を照合キーに揃える。使えない値は (None, 理由)。"""
    if value is None or str(value).strip() == '':
        return None, '空欄'
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = _clean(value)
    if re.fullmatch(r'B0[0-9A-Z]{8}', s):
        return None, 'ASINが入っている'
    if not s.isdigit():
        return None, '数字以外を含む'
    if len(s) not in (8, 12, 13):
        return None, f'{len(s)}桁'
    if not gtin_check_ok(s):
        return None, 'チェックディジット不一致'
    return ('0' + s if len(s) == 12 else s), ''


# ─────────────────────────────── 商品マスター ───────────────────────────────

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def _header_pos(header, wanted, where):
    pos = {}
    for k, h in wanted.items():
        if h not in header:
            raise MasterError(f'{where} に見出し「{h}」がありません。正しい商品マスターか確認してください')
        pos[k] = header.index(h)
    return pos


class Master:
    """商品マスターの読み取り専用索引。"""

    def __init__(self, path, sha256, products, by_jan_single, by_jan_set, jan_issues):
        self.path = path
        self.sha256 = sha256
        self.version = sha256[:12]
        self.products = products          # pid -> dict
        self.by_jan_single = by_jan_single  # 正規化JAN -> [pid](単品)
        self.by_jan_set = by_jan_set        # 正規化JAN -> [pid](セット)
        self.jan_issues = jan_issues        # [(pid, 元の値, 理由)] JANとして使えない行

    def resolve_jan(self, code):
        """正規化済みJANを照合する。推測はしない(重複は候補を返すだけ)。"""
        singles = self.by_jan_single.get(code, [])
        if len(singles) == 1:
            return {'status': R_MATCH, 'pid': singles[0], 'candidates': singles}
        if len(singles) > 1:
            return {'status': R_DUPLICATE, 'pid': '', 'candidates': singles}
        sets = self.by_jan_set.get(code, [])
        if sets:
            return {'status': R_SET_ONLY, 'pid': sets[0] if len(sets) == 1 else '', 'candidates': sets}
        return {'status': R_UNREGISTERED, 'pid': '', 'candidates': []}

    def describe(self, pid):
        p = self.products.get(pid)
        if not p:
            return None
        return {'pid': pid, 'name': p['name'], 'jan': p['jan_raw'], 'kind': p['kind'], 'cost': p['cost']}


def load_master(path=MASTER_FILE):
    """商品マスターを読み、P番号の一意性を確かめてから索引を作る。問題があれば止める。"""
    real = os.path.realpath(path)
    if real == os.path.realpath(LEGACY_SYSTEM_FILE) or os.path.basename(real) == os.path.basename(LEGACY_SYSTEM_FILE):
        raise MasterError('旧 在庫管理システム_v1.0.xlsm はID体系が異なるため照合に使えません')
    if not os.path.isfile(real):
        raise MasterError(f'商品マスターが見つかりません: {path}')
    digest = sha256_file(real)
    wb = load_workbook(real, read_only=True, data_only=True)
    try:
        if SH_MASTER not in wb.sheetnames:
            raise MasterError(f'シート「{SH_MASTER}」がありません: {path}')
        ws = wb[SH_MASTER]
        rows = ws.iter_rows(values_only=True)
        header = [str(c or '').strip() for c in next(rows)]
        pos = _header_pos(header, MASTER_HEAD, SH_MASTER)
        products, errors = {}, []
        for i, r in enumerate(rows, start=2):
            pid = str(r[pos['pid']] or '').strip()
            if not pid:
                if any(v not in (None, '') for v in r):
                    errors.append(f'行{i}: 内部管理IDが空なのにデータがある')
                continue
            if not PID_RE.match(pid):
                errors.append(f'行{i}: 内部管理IDの形式が不正 {pid!r}')
                continue
            if pid in products:
                errors.append(f'行{i}: 内部管理ID {pid} が重複(行{products[pid]["row"]})')
                continue
            cost = r[pos['cost']]
            products[pid] = {'row': i, 'jan_raw': r[pos['jan']], 'name': str(r[pos['name']] or '').strip(),
                             'kind': str(r[pos['kind']] or '').strip(),
                             'cost': cost if isinstance(cost, (int, float)) else None}
    finally:
        wb.close()
    if errors:
        # P番号の意味が一意でない状態で照合すると、別商品へ数量が入る。1件でもあれば止める
        raise MasterError('商品マスターのP番号に問題があります:\n  ' + '\n  '.join(errors[:20]))

    by_single, by_set, issues = defaultdict(list), defaultdict(list), []
    for pid, p in products.items():
        key, why = master_jan_key(p['jan_raw'])
        p['jan'] = key or ''
        if not key:
            issues.append((pid, p['jan_raw'], why))
            continue
        if p['kind'] not in (KIND_SINGLE, KIND_SET):
            # 単品かセットか分からない商品を単品として数えると換算を誤る → 照合対象にしない
            issues.append((pid, p['jan_raw'], f'商品種別が不明 {p["kind"]!r}'))
            continue
        (by_single if p['kind'] == KIND_SINGLE else by_set)[key].append(pid)
    return Master(real, digest, products, dict(by_single), dict(by_set), issues)
