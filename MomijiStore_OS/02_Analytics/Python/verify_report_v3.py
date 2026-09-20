# -*- coding: utf-8 -*-
"""verify_report_v3.py — 2026-08 経営レポート Ver.3 の数値を本番ファイルから再読込して照合する(読み取り専用)

照合するもの: KPI数値(楽天/Amazonチャネル/Amazon商品別/合計)・比率・在庫金額(サマリー)・在庫の内訳・原価確認3件の反映・
Amazon実額/商品別の区別(Q:R と N列が別の値であること)・注記の根拠(手数料15%行数・未配賦・梱包資材0・ナノックスワン行)
"""
import os
import sys
import warnings

warnings.filterwarnings('ignore')
from openpyxl import load_workbook

BASE = '/Users/hide0726/Desktop/Claude Code/MomijiStore_OS'
SD = BASE + '/02_Analytics/SourceData'
INV = BASE + '/01_InventoryManagement/SourceData'
REPORT = INV + '/Output/2026-08_経営レポート_Ver3.md'
fails = []


def chk(cond, msg):
    print(('  ✅ ' if cond else '  ❌ ') + msg)
    if not cond:
        fails.append(msg)


def block(ws):
    r = 8
    while ws.cell(r, 3).value is not None:
        r += 1
    labels = {}
    for rr in range(r + 1, ws.max_row + 1):
        lab = ws.cell(rr, 13).value
        if lab and str(lab).strip() not in labels:
            labels[str(lab).strip()] = rr
    return r, labels


rk = load_workbook(SD + '/楽天運営 KPI管理シート.xlsx', data_only=True)['8月']
am = load_workbook(SD + '/Amazon運営 KPI管理シート.xlsx', data_only=True)['8月']
amf = load_workbook(SD + '/Amazon運営 KPI管理シート.xlsx')['8月']
tr, rl = block(rk)
ta, al = block(am)
R = dict(sales=rk.cell(tr, 9).value, fee=rk.cell(tr, 11).value, cost=rk.cell(tr, 13).value, gp=rk.cell(tr, 14).value,
         margin=rk.cell(rl['限界利益'], 14).value, rate=rk.cell(rl['限界利益'] + 1, 14).value)
A = dict(sales=am.cell(ta, 9).value, fee_k=am.cell(ta, 11).value, cost=am.cell(ta, 13).value, gp_prod=am.cell(ta, 14).value,
         margin_prod=am.cell(al['限界利益'], 14).value)
# Q:R のチャネル全体KPI
qr = {str(am.cell(r, 17).value or ''): am.cell(r, 18).value for r in range(ta + 1, ta + 12) if am.cell(r, 17).value}
A['fee_ch'] = next(v for k, v in qr.items() if k.startswith('手数料 実額総額'))
A['gp_ch'] = next(v for k, v in qr.items() if k.startswith('粗利(チャネル'))
A['margin_ch'] = next(v for k, v in qr.items() if k.startswith('限界利益(チャネル'))
A['rate_ch'] = next(v for k, v in qr.items() if k.startswith('限界利益率(チャネル'))

print('■ KPI(本番ファイルの再読込)')
chk(round(R['sales']) == 2579795 and round(R['gp']) == 706377 and round(R['margin']) == 353895, f'楽天 売上 {R["sales"]:,} / 粗利 {R["gp"]:,.0f} / 限界利益 {R["margin"]:,.0f}')
chk(abs(R['gp'] / R['sales'] - 0.274) < 0.0005 and abs(R['rate'] - 0.137) < 0.0005, f'楽天 粗利率 {R["gp"]/R["sales"]:.1%} / 限界利益率 {R["rate"]:.1%}')
chk(round(R['fee']) == 386969 and round(R['fee']) == round(R['sales'] * 0.15), f'楽天 手数料 {R["fee"]:,.0f} = 売上×15%(推定)')
chk(round(A['sales']) == 5610675 and A['fee_ch'] == 759756 and round(A['gp_ch']) == 1669256 and round(A['margin_ch']) == 903493,
    f'Amazon チャネル: 売上 {A["sales"]:,} / 手数料実額 {A["fee_ch"]:,} / 粗利 {A["gp_ch"]:,.0f} / 限界利益 {A["margin_ch"]:,.0f}')
chk(abs(A['gp_ch'] / A['sales'] - 0.298) < 0.0005 and abs(A['rate_ch'] - 0.161) < 0.0005, f'Amazon チャネル 粗利率 {A["gp_ch"]/A["sales"]:.1%} / 限界利益率 {A["rate_ch"]:.1%}')
chk(round(A['gp_prod']) == 1649976 and round(A['margin_prod']) == 884213 and round(A['fee_k']) == 779036,
    f'Amazon 商品別(分離): 手数料K計 {A["fee_k"]:,.0f} / 粗利 {A["gp_prod"]:,.0f} / 限界利益 {A["margin_prod"]:,.0f}')
chk(round(A['gp_ch'] - A['gp_prod']) == 19280 and round(A['fee_k'] - A['fee_ch']) == 19280, 'Amazon 実額/商品別の差 ¥19,280 で整合')
tot_sales, tot_gp, tot_margin = R['sales'] + A['sales'], R['gp'] + A['gp_ch'], R['margin'] + A['margin_ch']
chk(round(tot_sales) == 8190470 and round(tot_gp) == 2375633 and round(tot_margin) == 1257388, f'合計 売上 {tot_sales:,} / 粗利 {tot_gp:,.0f} / 限界利益 {tot_margin:,.0f}')
chk(abs(tot_gp / tot_sales - 0.290) < 0.0005 and abs(tot_margin / tot_sales - 0.154) < 0.0005, f'合計 粗利率 {tot_gp/tot_sales:.1%} / 限界利益率 {tot_margin/tot_sales:.1%}')
chk(round(R['cost']) == 1486449 and round(A['cost']) == 3181663, f'売上原価 楽天 {R["cost"]:,.0f} / Amazon {A["cost"]:,.0f}(在庫の参考比率の分母)')
# 原価確認3件
def lv(ws, ch, ident):
    for r in range(8, ws.max_row + 1):
        if ws.cell(r, 3).value is None:
            break
        if str(ws.cell(r, 3).value).strip().upper() == ident.upper():
            return ws.cell(r, 12).value
chk(lv(rk, '楽天', 'b0915ngd2j-6') == 7536 and lv(am, 'Amazon', 'B0HD7N75QV') == 2862 and lv(am, 'Amazon', 'B0HB36RZPY') == 1232, 'F-01〜03 の原価が反映されている(7,536 / 2,862 / 1,232)')
# 注記の根拠
n15 = sum(1 for r in range(8, ta) if isinstance(amf.cell(r, 11).value, str) and amf.cell(r, 11).value.startswith('='))
chk(n15 == 432, f'Amazon 商品別の15%推定行 {n15}(未配賦53SKU ¥90,017 の注記の根拠)')
chk(am.cell(al['梱包資材'], 14).value == 0 and rk.cell(rl['クーポン利用手数料'], 14).value == 0, 'Amazon 梱包資材 0 / 楽天 クーポン利用手数料 0(確認済み)')
nano = [r for r in range(8, tr) if str(rk.cell(r, 3).value).lower() == 'b0cksbjbrp']
chk(len(nano) == 1 and rk.cell(nano[0], 9).value == 0 and rk.cell(nano[0], 8).value == 1 and rk.cell(nano[0], 12).value == 1200 and round(rk.cell(nano[0], 14).value) == -1200,
    f'ナノックスワン 行{nano[0] if nano else "?"}: 売上0・数量1・原価1,200・粗利−1,200')

print('■ 在庫(全体在庫サマリー・在庫リスト)')
sm = load_workbook(INV + '/全体在庫サマリー_v1.0.xlsx', data_only=True)['全体サマリー']
rows = {sm.cell(r, 1).value: (sm.cell(r, 2).value, sm.cell(r, 3).value, sm.cell(r, 4).value) for r in (5, 6, 7)}
chk(rows['楽天 単独在庫'] == (445, 3371, 7024963), f'楽天単独 {rows["楽天 単独在庫"]}')
chk(rows['Amazon FBA'] == (5, 14, 13952), f'FBA {rows["Amazon FBA"]}')
chk(rows['Amazon 自己発送'] == (581, 9497, 12852158), f'自己発送 {rows["Amazon 自己発送"]}')
tot = tuple(sum(v[i] for v in rows.values()) for i in range(3))
chk(tot == (1031, 12882, 19891073), f'合計 {tot}(システム在庫評価額(暫定) ¥19,891,073)')
notes = ' '.join(str(sm.cell(r, 1).value or '') for r in range(1, sm.max_row + 1))
chk('在庫基準日時 未取得' in notes and '確定在庫数量・確定評価額ではない' in notes, 'サマリーに「Amazon在庫基準日時 未取得」「確定在庫ではない」の注記')
solo = load_workbook(INV + '/Import/楽天在庫リスト_import.xlsx', read_only=True, data_only=True)['楽天単独在庫']
srows = [r for r in solo.iter_rows(min_row=2, values_only=True) if r and r[1]]
chk(len(srows) == 445 and not any(str(r[1]).lower() in ('b0842r98vl', 'b0dffzzxqw') for r in srows), '楽天単独在庫 445行・共有例外2件を含まない')
ratio = 19891073 / (R['cost'] + A['cost'])
chk(abs(ratio - 4.3) < 0.05, f'参考比率 在庫評価額÷8月売上原価 = {ratio:.1f} か月分(回転率ではない)')

print('■ レポート本文の記載')
txt = open(REPORT, encoding='utf-8').read()
for kw in ('システム在庫評価額(暫定)', '売上×15%推定', '759,756', '53SKU ¥90,017', '一次資料未照合', 'RMS原本どおり', '実地棚卸未実施', '梱包資材費は主として楽天側', '管理締め完了',
           '110行 → 最終確認16行 → 保留0件', '実地棚卸 → 在庫精度向上 → 在庫回転率 → 滞留在庫削減', '−¥753/セット'):
    chk(kw in txt, f'本文に「{kw}」')
for bad in ('会計上完全確定', '商品別利益完全確定', '実地棚卸済み', '確定在庫額', 'AIKOS'):
    chk(bad not in txt.replace('実地棚卸済み、のいずれでもありません', '').replace('商品別利益の完全確定・実地棚卸済み', ''), f'本文に禁止表現「{bad}」が無い')
print('\n' + ('PASS' if not fails else f'FAIL ({len(fails)}件)'))
sys.exit(0 if not fails else 1)
