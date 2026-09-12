# -*- coding: utf-8 -*-
"""cost_confirmations.py — 「人が確認した1販売分の原価」の記録(2026-09-12 ChatGPT承認 §7-6)

正本: 02_Analytics/SourceData/原価確認記録.xlsx(シート「記録」)。1行 = 1確認。

使い方:
    python3 cost_confirmations.py list [チャネル]
    python3 cost_confirmations.py add --ch 楽天 --ctrl b0c656jxkd --sku b0c656jxkd \\
        --from 2026-09 --to 2026-12 --cost 2283 --pack 3 --unit 単品原価 --pn "8507/2283-a" \\
        --scope "付属品なし" --by 英樹 --basis "2026-09-12 確認リストQ4-5 回答A(8月KPI L=2,283)" --ref-month 8月

なぜ必要か:
    KPIシートのL列に数値があるだけでは「人が確認した原価」とは言えない(自動で引き継いだ値かもしれない)。
    セルに数値を入れただけで自動的に確認済みにはしない。**確認者・確認日・対象出品・対象月・1販売分の原価・根拠**を
    1行ずつ残し、その記録と原価値を対応させる。既存の手入力原価を一括で確認済みにはしない。

記録が「いま使える」条件(§7-5):
    ・対象出品(チャネル・管理番号・SKU/ASIN)が一致する
    ・適用期間(対象月 開始〜終了)が生成月を含む。**確認日だけで無期限に引き継がない**
    ・販売構成(確認時の商品番号)が現在の商品番号と同じ(番号から読める構成・入数が一致)。違えば無効
    ・参照月を書いた記録は、その月のシートのL列が記録の原価と一致する。違えば無効(原価値が変わった)
    ・出品テーブルに販売入数があれば記録の入数と一致する。違えば無効
    ・無効化日が入っていない
"""
import argparse
import os
import shutil
import sys
from datetime import date

from openpyxl import Workbook, load_workbook

BASE = '/Users/hide0726/Desktop/Claude Code/MomijiStore_OS'
FILE = os.environ.get('COST_CONFIRM_FILE', BASE + '/02_Analytics/SourceData/原価確認記録.xlsx')
SHEET = '記録'
HEAD = ['記録ID', '確認日', '確認者', 'チャネル', '楽天商品管理番号', 'SKU/ASIN', '適用開始月', '適用終了月',
        '1販売分の原価', '販売入数', '原価単位', '含有範囲', '販売構成(確認時の商品番号)', '参照月', '根拠',
        '無効化日', '無効化理由']


def _norm(v):
    return str(v).strip().upper() if v is not None and str(v).strip() else ''


def load(ch=None):
    """記録を dict のリストで返す。ファイルが無ければ空(それが正しい状態)。"""
    if not os.path.exists(FILE):
        return []
    wb = load_workbook(FILE, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    hdr = [str(h or '').strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or r[0] in (None, ''):
            continue
        d = {h: (r[i] if i < len(r) else None) for i, h in enumerate(hdr)}
        if ch and str(d.get('チャネル') or '') != ch:
            continue
        out.append(d)
    return out


def _ym(v):
    """'2026-09' / '2026年9月' / datetime → '2026-09'。読めなければ ''。"""
    import re
    if v is None:
        return ''
    if hasattr(v, 'year'):
        return f'{v.year:04d}-{v.month:02d}'
    m = re.search(r'(\d{4})\D+(\d{1,2})', str(v))
    return f'{int(m.group(1)):04d}-{int(m.group(2)):02d}' if m else ''


def covers(rec, ym):
    a, b = _ym(rec.get('適用開始月')), _ym(rec.get('適用終了月'))
    if not a:
        return False
    if not b:
        b = a                                   # 終了月が無ければ単月
    return a <= ym <= b


def find(records, ch, ctrl, sku, ym):
    """使える記録を新しい確認日順に返す(適用期間と出品が合うものだけ。構成・原価の照合は呼び出し側)。"""
    hits = [r for r in records
            if str(r.get('チャネル') or '') == ch and _norm(r.get('楽天商品管理番号')) == _norm(ctrl)
            and _norm(r.get('SKU/ASIN')) == _norm(sku) and covers(r, ym) and not r.get('無効化日')]
    hits.sort(key=lambda r: str(r.get('確認日') or ''), reverse=True)
    return hits


def add(rec):
    """1行追加する(承認後にだけ呼ぶ)。既存ファイルはバックアップしてから書く。"""
    os.makedirs(os.path.dirname(FILE), exist_ok=True)
    if os.path.exists(FILE):
        bdir = os.path.dirname(FILE) + '/Backup'
        os.makedirs(bdir, exist_ok=True)
        shutil.copy2(FILE, f'{bdir}/原価確認記録_backup_{date.today():%Y%m%d_%H%M%S}.xlsx')
        wb = load_workbook(FILE)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.create_sheet(SHEET)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
    if ws.max_row < 1 or ws.cell(1, 1).value != HEAD[0]:
        for c, h in enumerate(HEAD, 1):
            ws.cell(1, c).value = h
    n = sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0] not in (None, ''))
    rec = dict(rec)
    rec['記録ID'] = f'K{n + 1:05d}'
    rec.setdefault('確認日', date.today().isoformat())
    row = ws.max_row + 1 if n else 2
    for c, h in enumerate(HEAD, 1):
        ws.cell(row, c).value = rec.get(h)
    # 識別子は文字列のまま保つ(先頭0を落とさない)
    for c in (5, 6, 13):
        ws.cell(row, c).number_format = '@'
    wb.save(FILE)
    return rec['記録ID']


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest='cmd')
    l = sub.add_parser('list'); l.add_argument('ch', nargs='?')
    a = sub.add_parser('add')
    for k in ('ch', 'ctrl', 'sku', 'cost', 'by', 'basis', 'pn', 'scope'):
        a.add_argument('--' + k, required=True)
    a.add_argument('--from', dest='from_', required=True)
    a.add_argument('--to', default='')
    a.add_argument('--pack', default='')
    a.add_argument('--unit', default='')
    a.add_argument('--ref-month', default='')
    args = ap.parse_args()
    if args.cmd == 'list':
        recs = load(args.ch)
        print(f'{len(recs)}件  ({FILE})')
        for r in recs:
            print('  ', {k: r.get(k) for k in ('記録ID', '確認日', 'チャネル', '楽天商品管理番号', 'SKU/ASIN', '適用開始月', '適用終了月', '1販売分の原価', '無効化日')})
    elif args.cmd == 'add':
        rid = add({'確認者': args.by, 'チャネル': args.ch, '楽天商品管理番号': args.ctrl, 'SKU/ASIN': args.sku,
                   '適用開始月': args.from_, '適用終了月': args.to, '1販売分の原価': float(args.cost),
                   '販売入数': (int(args.pack) if args.pack else None), '原価単位': args.unit, '含有範囲': args.scope,
                   '販売構成(確認時の商品番号)': args.pn, '参照月': args.ref_month, '根拠': args.basis})
        print(f'✅ 記録 {rid} を追加: {FILE}')
    else:
        print(__doc__)


if __name__ == '__main__':
    main()
