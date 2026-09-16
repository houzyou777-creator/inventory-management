# -*- coding: utf-8 -*-
"""verify_v4_copy.py — V-4(集計先の書式を書き込むセルだけにする)をコピーで検証する(読み取り専用)

使い方:
    python3 verify_v4_copy.py 1      1回目(Import = 8/3資料 864行)の後
    python3 verify_v4_copy.py 2      2回目(Import = 9/16資料 814行・行数が減る)の後
    python3 verify_v4_copy.py prod   本番へ V-4 を入れて保存しただけの状態(データ不変)を確認する

確認すること(ChatGPT 2026-09-16 Q17 の条件):
    1. 保存・閉じる・再読込後のファイルサイズ(前後比。V-3 は 307KB → 6,985KB になった)
    2. 識別子の完全一致(取込 = Import の文字列、集計/要確認 = 取込の文字列)、数値化0
    3. 空欄保持(取込・集計・要確認の識別子で、元が空欄なら本当の空欄。長さ0の文字列にしない)
    4. 集計値(総在庫金額・数量・未登録)が期待値と一致
    5. 数式・見出しの書式・ボタン(図形)・VBA が保たれている
    6. 2回目で行数が減っても旧データが残らない(集計・要確認・取込のデータ行より下が空)

何も書き換えない。判定は PASS / FAIL を最後に出す。
"""
import os
import sys
import warnings
import zipfile

warnings.filterwarnings('ignore')
from openpyxl import load_workbook

SD = '/Users/hide0726/Desktop/Claude Code/MomijiStore_OS/01_InventoryManagement/SourceData'
PROD = SD + '/楽天在庫金額集計ツール_v1.0.xlsm'
BASELINE = SD + '/Backup/楽天在庫金額集計ツール_v1.0_backup_20260916_0612時点(既にV3差し替え済み).xlsm'   # V-3 本番・集計前(307KB)
COPY = SD + '/V4_test/楽天在庫金額集計ツール_v1.0.xlsm'
IMPORT1 = SD + '/V4_test/Import/楽天在庫リスト_import.xlsx'
IMPORT2 = SD + '/V4_test/Import_2回目/楽天在庫リスト_import.xlsx'
BAS = os.path.dirname(SD) + '/VBA/Module_Rakuten_Tool.bas'
EXPECT = {1: dict(amount=15077081, qty=6801, unreg=20, rows=864),      # 8/3資料: Python予測(引継ぎ107)。V-3テストコピー実測 15,064,681/21 は原価マスター復元前(クリニーク12,400の差)
          2: dict(amount=12621422, qty=5911, unreg=59, rows=814)}      # 9/16資料: V3_stock_test コピーの Excel 実測(引継ぎ123)
SIZE_LIMIT = 1_000_000   # これを超えたら肥大(V-3 は 6,985,292。V-4 1回目(2026-09-16 17:52)は Interior.ColorIndex=xlNone の10万行で 4,714,370)

fails = []


def check(cond, msg):
    print(('  ✅ ' if cond else '  ❌ ') + msg)
    if not cond:
        fails.append(msg)


def cstr(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return v.strip() if isinstance(v, str) else str(v)


def data_rows(ws, min_row, key_cols, width=13):
    out = []
    for r in ws.iter_rows(min_row=min_row, values_only=True):
        r = tuple(r) + (None,) * (width - len(r))          # 右端の空列が省かれても添字が揃うように
        if any(r[c] not in (None, '') for c in key_cols):
            out.append(r)
    return out


def zip_sizes(path):
    z = zipfile.ZipFile(path)
    return {i.filename: i.file_size for i in z.infolist()}, os.path.getsize(path)


def vba_modules(path):
    from oletools.olevba import VBA_Parser
    vp = VBA_Parser(path)
    out = {name: code.replace('\r\n', '\n') for (_, _, name, code) in vp.extract_macros()}
    vp.close()
    return out


def body(s):
    # VBE(cp932)を往復すると「〜」が「～」になる(コメントのみ)。比較ではそろえる
    return '\n'.join(l for l in s.split('\n') if not l.startswith('Attribute VB_')).replace('\u301c', '\uff5e')


def check_vba(target, label):
    mods = vba_modules(target)
    base = vba_modules(BASELINE)
    bas = open(BAS, encoding='utf-8').read().replace('\r\n', '\n')
    check(body(mods.get('Module_Rakuten_Tool.bas', '')) == body(bas), f'{label}: Module_Rakuten_Tool = VBA/Module_Rakuten_Tool.bas(V-4)')
    others = [k for k in base if k != 'Module_Rakuten_Tool.bas']
    check(all(mods.get(k) == base[k] for k in others) and set(mods) == set(base), f'{label}: 他の {len(others)} モジュールは変更なし')
    v3 = body(base['Module_Rakuten_Tool.bas'])
    m = mods.get('Module_Rakuten_Tool.bas', '')
    check('ClearDataArea' in m and 'PutIdText' in m and 'Cells(100000, AG_JAN)' not in m and 'A10:L100000' not in m and 'A3:H100000' not in m,
          f'{label}: V-4 の中身(ClearDataArea・PutIdText あり／10万行への書式操作 "@"・Interior なし)')
    check('Cells(100000, AG_JAN)' in v3, '基準(V-3)には10万行の一括書式がある(比較の前提)')


def check_static(target, label):
    """数式・見出し書式・図形(ボタン)・シート構成が基準と同じか。"""
    zt, st = zip_sizes(target)
    zb, sb = zip_sizes(BASELINE)
    draw_t = {k: v for k, v in zt.items() if 'drawing' in k or 'vml' in k}
    draw_b = {k: v for k, v in zb.items() if 'drawing' in k or 'vml' in k}
    check(draw_t == draw_b, f'{label}: 図形・ボタンのパーツ({len(draw_b)}個)が基準と同じ大きさ {sorted(draw_b.items())}')
    wt = load_workbook(target)
    wb = load_workbook(BASELINE)
    check(wt.sheetnames == wb.sheetnames, f'{label}: シート構成 {wt.sheetnames}')
    # 数式(全シート)
    ft = {(s, c.coordinate): c.value for s in wt.sheetnames for row in wt[s].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')}
    fb = {(s, c.coordinate): c.value for s in wb.sheetnames for row in wb[s].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')}
    check(ft == fb, f'{label}: 数式 {len(fb)} 個が基準と同じ')
    # 見出し行の書式(集計 1〜9行・要確認 1〜2行・取込 1〜2行・原価マスター 1〜2行)
    def fmt(ws, rows, cols):
        return [(r, c, ws.cell(r, c).number_format, ws.cell(r, c).font.name, ws.cell(r, c).font.b, ws.cell(r, c).fill.fgColor.rgb, ws.cell(r, c).border.bottom.style)
                for r in rows for c in cols]
    for sh, rows, cols in (('在庫金額集計', range(1, 10), range(1, 13)), ('要確認一覧', range(1, 3), range(1, 9)),
                           ('楽天CSV取込', range(1, 3), range(1, 13)), ('原価マスター', range(1, 3), range(1, 8))):
        check(fmt(wt[sh], rows, cols) == fmt(wb[sh], rows, cols), f'{label}: {sh} 見出し行の書式が基準と同じ')
    # 金額列の書式(集計 F/H/I/J のデータ行)は #,##0 のまま
    ws = wt['在庫金額集計']
    n = len(data_rows(ws, 10, (0, 1, 2)))
    if n:
        check(all(ws.cell(r, c).number_format == '#,##0' for r in range(10, 10 + n) for c in (6, 8, 9, 10) if ws.cell(r, c).value not in (None, '')),
              f'{label}: 集計 金額列(F/H/I/J)の書式 #,##0 が {n} 行で保たれている')
    # 原価マスターのデータ不変
    mt = [tuple(r) for r in wt['原価マスター'].iter_rows(values_only=True)]
    mb = [tuple(r) for r in wb['原価マスター'].iter_rows(values_only=True)]
    check(mt == mb, f'{label}: 原価マスター 全セル不変({len(mb)}行)')
    return st, sb


def check_run(run):
    imp = IMPORT1 if run == 1 else IMPORT2
    exp = EXPECT[run]
    print(f'■ V-4 コピー検証 {run}回目: {os.path.basename(COPY)} / Import={os.path.relpath(imp, SD)}')
    # 1. サイズ
    st, sb = check_static(COPY, f'{run}回目')
    zt, _ = zip_sizes(COPY)
    print(f'   サイズ: 基準 {sb:,} → コピー {st:,} bytes / sheet4(集計) {zt.get("xl/worksheets/sheet4.xml", 0):,} / sheet5(要確認) {zt.get("xl/worksheets/sheet5.xml", 0):,}')
    check(st < SIZE_LIMIT, f'{run}回目: 保存後のファイルが {st:,} bytes(< {SIZE_LIMIT:,}。V-3 は 6,985,292・V-4 1回目 4,714,370)')
    import re as _re
    z = zipfile.ZipFile(COPY)
    nrow4 = z.read('xl/worksheets/sheet4.xml').count(b'<row ')
    nrow5 = z.read('xl/worksheets/sheet5.xml').count(b'<row ')
    check(nrow4 < exp['rows'] + 300 and nrow5 < 1000, f'{run}回目: 保存された行要素 集計 {nrow4:,}・要確認 {nrow5:,}(10万行ではない)')
    check_vba(COPY, f'{run}回目')
    # 2. 取込 = Import
    wi = load_workbook(imp, read_only=True, data_only=True)['在庫']
    irows = data_rows(wi, 2, (1, 3))
    wc = load_workbook(COPY, data_only=True)
    tk = wc['楽天CSV取込']
    trows = data_rows(tk, 3, (1, 3))
    check(len(trows) == len(irows) == exp['rows'], f'{run}回目: 取込 {len(trows)} 行 = Import {len(irows)} 行 = 期待 {exp["rows"]}')
    mism = typebad = blankbad = 0
    for a, b in zip(irows, trows):
        for c in range(4):
            e, g = cstr(a[c]), b[c]
            if e is None or e == '':
                blankbad += g not in (None,)
            else:
                mism += (g != e)
                typebad += not isinstance(g, str)
    check(mism == 0 and typebad == 0, f'{run}回目: 取込の識別子4列が Import と完全一致(不一致 {mism}・非文字列 {typebad})')
    check(blankbad == 0, f'{run}回目: 取込の空欄は本当の空欄(長さ0文字列 {blankbad})')
    idchk = sum(1 for r in trows if r[11] not in (None, ''))
    check(idchk == 0, f'{run}回目: 識別子チェック(L列)の記録 {idchk} 件')
    # データ行より下が空(取込)
    below = [r for r in tk.iter_rows(min_row=3 + len(trows), values_only=True) if r and any(v not in (None, '') for v in r[:12])]
    check(len(below) == 0, f'{run}回目: 取込のデータ行より下に旧データなし({len(below)})')
    # 3. 集計
    ag = wc['在庫金額集計']
    ag_dims, ag_max = ag.dimensions, ag.max_row          # ws.cell() は無いセルを作って使用範囲を広げるので先に取る
    amount, qty, unreg = ag.cell(3, 2).value, ag.cell(4, 2).value, ag.cell(5, 6).value
    print(f'   集計: 総在庫金額 {amount:,} / 総在庫数量 {qty:,} / 未登録 {unreg} / 集計日時 {ag.cell(6, 2).value}')
    check(round(amount) == exp['amount'] and qty == exp['qty'] and unreg == exp['unreg'],
          f'{run}回目: 集計値 = 期待値({exp["amount"]:,} / {exp["qty"]:,} / {exp["unreg"]})')
    arows = data_rows(ag, 10, (0, 1, 2))
    check(len(arows) == exp['rows'], f'{run}回目: 集計 {len(arows)} 行 = 期待 {exp["rows"]}')
    below = [r for r in ag.iter_rows(min_row=10 + len(arows), values_only=True) if r and any(v not in (None, '') for v in r[:12])]
    check(len(below) == 0, f'{run}回目: 集計のデータ行より下に旧データなし({len(below)})')
    # 集計の識別子 = 取込の識別子(集合として)、文字列、空欄は None
    tset = {}
    for r in trows:
        tset[(r[1] or '', r[3] or '')] = r
    numeric = sum(1 for r in arows for c in (0, 1, 2, 11) if isinstance(r[c], (int, float)))
    zerolen = sum(1 for r in arows for c in (0, 1, 2, 11) if r[c] == '')
    check(numeric == 0, f'{run}回目: 集計の識別子(JAN/管理番号/SKU/商品番号)に数値型なし({numeric})')
    check(zerolen == 0, f'{run}回目: 集計の識別子に長さ0の文字列なし({zerolen})。空欄は空欄')
    diff = 0
    for r in arows:
        t = tset.get((r[1] or '', r[2] or ''))
        if t is None or cstr(t[0]) != r[0] or cstr(t[2]) != r[11]:
            diff += 1
    check(diff == 0, f'{run}回目: 集計の JAN・商品番号が取込と一致(不一致 {diff})')
    fmt_bad = sum(1 for i in range(10, 10 + len(arows)) for c in (1, 2, 3, 12)
                  if ag.cell(i, c).value not in (None, '') and ag.cell(i, c).number_format != '@')
    check(fmt_bad == 0, f'{run}回目: 集計の書き込んだ識別子セルは "@"(違反 {fmt_bad})')
    tail_fmt = sum(1 for i in (10 + len(arows), 10 + len(arows) + 50, 5000, 99999) for c in (1, 2, 3, 12) if ag.cell(i, c).number_format == '@')
    check(tail_fmt == 0, f'{run}回目: 集計のデータ行より下の識別子列に "@" が残っていない({tail_fmt})')
    check(ag_max < 10 + len(arows) + 200, f'{run}回目: 集計シートの使用範囲 {ag_dims}(10万行になっていない)')
    # 要確認
    ck = wc['要確認一覧']
    ck_dims, ck_max = ck.dimensions, ck.max_row
    crows = data_rows(ck, 3, (0, 2))
    reasons = {}
    for r in crows:
        reasons[r[4]] = reasons.get(r[4], 0) + 1
    print(f'   要確認一覧 {len(crows)} 行: {reasons}')
    check(reasons.get('原価未登録', 0) == exp['unreg'], f'{run}回目: 要確認一覧の「原価未登録」{reasons.get("原価未登録", 0)} 行 = 未登録 {exp["unreg"]}')
    below = [r for r in ck.iter_rows(min_row=3 + len(crows), values_only=True) if r and any(v not in (None, '') for v in r[:8])]
    check(len(below) == 0, f'{run}回目: 要確認一覧のデータ行より下に旧データなし({len(below)})')
    zerolen = sum(1 for r in crows for c in (0, 1, 2) if r[c] == '')
    numeric = sum(1 for r in crows for c in (0, 1, 2) if isinstance(r[c], (int, float)))
    check(zerolen == 0 and numeric == 0, f'{run}回目: 要確認一覧の識別子: 長さ0文字列 {zerolen}・数値型 {numeric}(V-3 では B30 に長さ0があった)')
    check(ck_max < 3 + len(crows) + 200, f'{run}回目: 要確認一覧の使用範囲 {ck_dims}')


def check_prod():
    print('■ 本番に V-4 を入れて保存しただけの状態を確認(データ不変)')
    st, sb = check_static(PROD, '本番')
    print(f'   サイズ: 基準 {sb:,} → 本番 {st:,} bytes')
    check_vba(PROD, '本番')
    wt = load_workbook(PROD, data_only=True)
    wb = load_workbook(BASELINE, data_only=True)
    for sh in wb.sheetnames:
        a = [tuple(r) for r in wt[sh].iter_rows(values_only=True)]
        b = [tuple(r) for r in wb[sh].iter_rows(values_only=True)]
        check(a == b, f'本番: {sh} 全セル値が V-4 差し替え前と同じ({len(b)}行)')


if __name__ == '__main__':
    arg = sys.argv[1] if len(sys.argv) > 1 else '1'
    if arg == 'prod':
        check_prod()
    else:
        check_run(int(arg))
    print('\n' + ('PASS' if not fails else f'FAIL ({len(fails)}件)'))
    for f in fails:
        print('  -', f)
    sys.exit(0 if not fails else 1)
