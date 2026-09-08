# -*- coding: utf-8 -*-
"""apply_monthly_expenses.py — KPIシートの月次経費(黄色セル)を確認・入力する

使い方:
    python3 apply_monthly_expenses.py list <月>
    python3 apply_monthly_expenses.py set <月> <チャネル> <費目> <金額> [...]

例:
    python3 apply_monthly_expenses.py list 8月
    python3 apply_monthly_expenses.py set 8月 楽天 広告費 53095 Amazon 広告費 111379

なぜスクリプトにするか:
    経費行の位置は**毎月ずれる**。商品行数が変わると合計行も経費行も動く。
    「N272へ書く」と覚えると、翌月には別の費目へ書き込む事故になる。
    M列のラベル(広告費・送料など)を探して書くので、行がずれても正しく入る。

処理:
    1. KPIシートをBackup/へバックアップ
    2. M列のラベルを探して N列へ金額を書き、黄色塗りを外す
    3. Excel(AppleScript)で再計算し、限界利益が更新されたことを確認する

⚠️ 金額は人が決めた値だけを入れる。AIが推定した値は入れない。
"""
import os
import shutil
import subprocess
import sys
from datetime import date

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

BASE = '/Users/hide0726/Desktop/Claude Code/MomijiStore_OS'
SD = BASE + '/02_Analytics/SourceData'
FILES = {'楽天': SD + '/楽天運営 KPI管理シート.xlsx',
         'Amazon': SD + '/Amazon運営 KPI管理シート.xlsx'}
NO_FILL = PatternFill(fill_type=None)

# 費目ごとの「どこから取るか」「誰が入れるか」。②黄色セル一覧の出どころでもある
SOURCES = {
    '広告費': ('広告改善一覧 / 広告分析シートの広告費合計',
               'AI(分析から自動算出) → 人が承認'),
    'ポイント費用': ('RMS → 売上・実績 → ポイント実績', '人(RMSから転記)'),
    'クーポン利用額': ('RMS → 販促 → クーポン実績', '人(RMSから転記)'),
    'クーポン利用手数料': ('RMS → 請求明細', '人(RMSから転記)'),
    'プロモーション費': ('セラーセントラル → 支払い → トランザクション', '人(転記)'),
    'その他手数料': ('セラーセントラル → 支払い → トランザクション', '人(転記)'),
    '送料': ('ヤマト/佐川の月次請求書。AmazonはFBA配送料',
             '人(請求書から転記)。⚠️ 配送マスター完成後は自動算出へ'),
    '梱包資材': ('仕入先の月次請求書', '人(請求書から転記)'),
}


def find_block(ws):
    """合計行と経費ラベルの位置を**ラベルを探して**求める。行番号は決め打ちしない。"""
    r = 8
    while ws.cell(r, 3).value is not None:
        r += 1
    total_row = r                                  # 商品の合計行
    labels = {}
    for rr in range(total_row + 1, ws.max_row + 1):
        lab = ws.cell(rr, 13).value                # M列 = 費目名
        if lab and lab not in ('合計',) and lab not in labels:
            labels[str(lab).strip()] = rr
    return total_row, labels


def is_yellow(c):
    f = c.fill
    return bool(f and f.fgColor and f.fgColor.rgb in ('FFFFFF00', '00FFFF00'))


def survey(month):
    """黄色セル(未入力の経費)を洗い出す。仕入値の未入力行数も数える。"""
    out = []
    for ch, path in FILES.items():
        wb = load_workbook(path)
        if month not in wb.sheetnames:
            continue
        ws = wb[month]
        total_row, labels = find_block(ws)
        for lab, rr in labels.items():
            if lab in ('限界利益', '限界利益率'):     # 計算結果。入力対象ではない
                continue
            c = ws.cell(rr, 14)                    # N列 = 金額
            src, who = SOURCES.get(lab, ('—', '人'))
            out.append({'ch': ch, 'item': lab, 'cell': c.coordinate,
                        'value': c.value, 'filled': c.value is not None,
                        'yellow': is_yellow(c), 'src': src, 'who': who})
        # 仕入値(L列)の未入力は商品単位の話。件数だけ持つ
        n = sum(1 for r in range(8, total_row) if is_yellow(ws.cell(r, 12)))
        if n:
            out.append({'ch': ch, 'item': f'仕入値(商品ごと)', 'cell': f'L列 {n}行',
                        'value': None, 'filled': False, 'yellow': True,
                        'src': '商品マスター E列 / 仕入先の納品書',
                        'who': '人 → 要対応一覧 02・03 で対象商品を確認'})
    return out


def apply(month, pairs):
    """(チャネル, 費目, 金額) を書き込む。ラベルが無ければ止める。"""
    touched = {}
    for ch, item, amount in pairs:
        if ch not in FILES:
            sys.exit(f'❌ 知らないチャネルです: {ch}(楽天 / Amazon)')
        path = FILES[ch]
        wb = touched.get(ch) or load_workbook(path)
        touched[ch] = wb
        if month not in wb.sheetnames:
            sys.exit(f'❌ {ch} に {month} シートがありません')
        ws = wb[month]
        _, labels = find_block(ws)
        if item not in labels:
            sys.exit(f'❌ {ch} {month} に費目「{item}」がありません。'
                     f'あるのは: {" / ".join(labels)}')
        c = ws.cell(labels[item], 14)
        before = c.value
        c.value = amount
        c.fill = NO_FILL
        print(f'  {ch} {month} {item}: {c.coordinate} '
              f'{"(空)" if before is None else before} → {amount:,}')

    for ch, wb in touched.items():
        path = FILES[ch]
        bdir = os.path.dirname(path) + '/Backup'
        os.makedirs(bdir, exist_ok=True)
        base = os.path.splitext(os.path.basename(path))[0]
        shutil.copy2(path, f'{bdir}/{base}_backup_{date.today():%Y%m%d}_{month}経費入力前.xlsx')
        wb.save(path)
    return list(touched)


def recalc(path):
    script = f'''
set p to POSIX file "{path}"
with timeout of 600 seconds
tell application "Microsoft Excel"
    set wasRunning to running
    open p
    delay 2
    calculate
    save active workbook
    close active workbook saving no
    if not wasRunning then quit
end tell
end timeout
return "ok"
'''
    r = subprocess.run(['osascript', '-e', script], capture_output=True,
                       text=True, timeout=660)
    if r.returncode != 0:
        raise RuntimeError(f'Excel再計算に失敗: {r.stderr}')


def show_margin(ch, month):
    """限界利益と限界利益率を読む。未確定なら何が足りないかも返す。"""
    wb = load_workbook(FILES[ch], data_only=True)
    ws = wb[month]
    total_row, labels = find_block(ws)
    gp = ws.cell(total_row, 14).value
    sales = ws.cell(total_row, 9).value
    mp = labels.get('限界利益')
    margin = ws.cell(mp, 14).value if mp else None
    rate = ws.cell(mp + 1, 14).value if mp else None
    missing = [lab for lab, rr in labels.items()
               if lab not in ('限界利益', '限界利益率')
               and ws.cell(rr, 14).value is None]
    return sales, gp, margin, rate, missing


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    mode, month = sys.argv[1], sys.argv[2]

    if mode == 'list':
        rows = survey(month)
        print(f'■ {month} の月次経費(黄色セル)')
        print(f'  {"チャネル":8} {"費目":18} {"セル":10} {"状態":10} 取得元 / 入力者')
        for d in rows:
            st = '入力済' if d['filled'] else '🟡未入力'
            print(f'  {d["ch"]:8} {d["item"]:18} {d["cell"]:10} {st:10} '
                  f'{d["src"]} / {d["who"]}')
        for ch in FILES:
            s, gp, m, rt, miss = show_margin(ch, month)
            print(f'\n  {ch} {month}: 売上 ¥{s:,.0f} / 粗利 ¥{gp:,.0f}'
                  f' / 限界利益 {"¥{:,.0f}".format(m) if m else "—"}'
                  f'{f" ({rt:.1%})" if isinstance(rt, float) else ""}')
            if miss:
                print(f'    ⚠️ 未入力のため確定していない: {" / ".join(miss)}')
        return

    if mode != 'set' or len(sys.argv) < 6 or (len(sys.argv) - 3) % 3:
        print(__doc__)
        sys.exit(1)
    a = sys.argv[3:]
    pairs = [(a[i], a[i + 1], float(a[i + 2])) for i in range(0, len(a), 3)]
    pairs = [(c, i, int(v) if v == int(v) else v) for c, i, v in pairs]

    print(f'■ {month} の経費を入力します')
    chs = apply(month, pairs)
    print('Excelで再計算中...')
    for ch in chs:
        recalc(FILES[ch])
    print('\n■ 反映後')
    for ch in chs:
        s, gp, m, rt, miss = show_margin(ch, month)
        print(f'  {ch} {month}: 売上 ¥{s:,.0f} / 粗利 ¥{gp:,.0f}'
              f' / 限界利益 {"¥{:,.0f}".format(m) if m else "—"}'
              f'{f" ({rt:.1%})" if isinstance(rt, float) else ""}')
        if miss:
            print(f'    ⚠️ まだ未入力: {" / ".join(miss)} → 限界利益は未確定')
        else:
            print('    ✅ 経費がすべて入り、限界利益が確定しました')


if __name__ == '__main__':
    main()
