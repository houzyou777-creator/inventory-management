# -*- coding: utf-8 -*-
"""pricetar_cost_diff.py — プライスター在庫CSVの cost と 商品マスター・出品テーブル・原価確認記録 の差分レポート(Step A)

使い方:
    python3 pricetar_cost_diff.py <プライスター在庫CSV> [<CSV2> ...] [--month 2026-09] [--out <出力xlsx>]

例:
    python3 pricetar_cost_diff.py \\
        ../../01_InventoryManagement/import/Amazon/1296040_inventoryList_202608_2.csv \\
        ../../01_InventoryManagement/import/Amazon/1296040_inventoryList_FBA_202608_2.csv

**読み取り専用。** 商品マスター・出品テーブル・原価確認記録・KPIシートには一切書き込まない
(read_only で開き、保存しない)。書き出すのは新規の差分レポート xlsx 1本だけ。

なぜ必要か(2026-09-25 承認 Step A):
    プライスターの cost は「最新仕入値の有力な観測値」だが、商品マスターの正ではない。
    しかも cost が「単品1個分」か「セット全体」かは出品ごとに違いうる。
    そこで**決めつけずに**分類だけを行い、どちらが正しいかは人が承認欄で判断する。

分類(優先順):
    出品紐付け不可 / 紐付け曖昧 / プライスター原価なし / 商品マスター原価なし /
    一致 / 整数倍でセット商品の可能性あり / 原価差異あり

照合キー(Master_Design §4.2・既存ルールに従う):
    ・SKU を最優先(プライスターの SKU = Amazon の出品SKU)
    ・SKU で引けない時だけ ASIN。同じASINに複数出品があれば**推測せず「紐付け曖昧」**
    ・JAN は照合に使わない(重複・入数違いの同一JANがあるため)。表示のみ
"""
import argparse
import csv
import hashlib
import io
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import momiji_paths as MP   # 保存場所は momiji_paths が __file__ から決める(直書きしない)

MASTER_FILE = MP.INV_SD + '/商品マスター_単品_v1.0.xlsx'
CONFIRM_FILE = MP.ANA_SD + '/原価確認記録.xlsx'
OUTPUT_DIR = MP.INV_OUTPUT

SH_MASTER = '商品マスター'
SH_LISTING = '出品テーブル'
SH_CONFIRM = '記録'

# 列は位置ではなく見出し名で引く(列が足されても壊れないように)
MASTER_HEAD = {'pid': '内部管理ID', 'jan': 'JAN', 'name': '商品名', 'type': '商品種別', 'cost': '標準原価'}
LISTING_HEAD = {'lid': '出品ID', 'pid': '内部管理ID', 'ch': 'チャネル', 'asin': 'ASIN',
                'asku': 'AmazonSKU', 'pack': '販売入数', 'unit': '原価単位', 'proof': '確認根拠'}
CSV_NEED = ['SKU', 'ASIN', 'number', 'price', 'cost', 'title']

UNIT_SINGLE = '単品原価'
UNIT_SET = 'セット原価'

CAT_NOLINK = '出品紐付け不可'
CAT_AMBIG = '紐付け曖昧'
CAT_NO_PC = 'プライスター原価なし'
CAT_NO_MC = '商品マスター原価なし'
CAT_MATCH = '一致'
CAT_MULTI = '整数倍でセット商品の可能性あり'
CAT_DIFF = '原価差異あり'
# レポートの並び順(人が見るべきものを先に)
CAT_ORDER = [CAT_DIFF, CAT_MULTI, CAT_AMBIG, CAT_NOLINK, CAT_NO_MC, CAT_NO_PC, CAT_MATCH]

RATIO_TOL = 0.01        # 整数倍とみなす許容(倍率の ±1%)。参考判定であり確定ではない
SMALL_DIFF = 10         # 10円以下は少額差(既存 push-kpi と同じ基準)

ADVICE = {
    CAT_MATCH: '対応不要',
    CAT_DIFF: '仕入明細で最新仕入値を確認し、正しい値を判断(マスター更新 or プライスター修正)。承認欄に記入',
    CAT_MULTI: '出品の販売入数・原価単位(単品原価/セット原価)を確認(Step B 候補)。マスター原価は変更しない',
    CAT_NO_MC: 'プライスターcostを候補値として、単位(単品/セット)を確認のうえ標準原価の登録を判断',
    CAT_NO_PC: 'プライスター側の仕入値未入力。入力要否を判断(入力は人がプライスター上で行う)',
    CAT_NOLINK: '出品テーブルへの登録要否を確認(新規出品・SKU変更の可能性)',
    CAT_AMBIG: '候補出品の中から正しい出品を特定し、出品テーブルのSKUを確認',
}


# ─────────────────────────────── 読み込み ───────────────────────────────

def norm(v):
    return str(v).strip().upper() if v is not None and str(v).strip() else ''


def id_str(v):
    """識別子を文字列で返す(JAN が数値で入っていても先頭0以外を崩さない)。"""
    if v is None or str(v).strip() == '':
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def to_num(v):
    try:
        s = str(v).replace(',', '').strip()
        return float(s) if s != '' else None
    except (TypeError, ValueError):
        return None


def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def read_pricetar_csv(path):
    """プライスター在庫CSV(Shift_JIS または UTF-8 BOM・値は ="..." 形式)を dict のリストで返す。"""
    with open(path, 'rb') as f:
        raw = f.read()
    txt = raw.decode('utf-8-sig') if raw.startswith(b'\xef\xbb\xbf') else raw.decode('cp932')
    reader = csv.DictReader(io.StringIO(txt))
    missing = [c for c in CSV_NEED if c not in (reader.fieldnames or [])]
    if missing:
        # 列仕様が変わったら黙って続けない(誤った列を原価として読む事故を防ぐ)
        raise SystemExit(f'CSVの列が想定と違います({os.path.basename(path)}): 不足 {missing}')
    out = []
    for i, r in enumerate(reader, start=2):
        rec = {k: re.sub(r'^="?|"$', '', (v or '').strip()) for k, v in r.items() if k}
        rec['_file'] = os.path.basename(path)
        rec['_line'] = i
        out.append(rec)
    return out


def _header_pos(ws, wanted, sheet):
    header = [str(c or '').strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    pos = {}
    for k, h in wanted.items():
        if h not in header:
            raise SystemExit(f'{sheet} に見出し「{h}」がありません。列構成を確認してください')
        pos[k] = header.index(h)
    return pos


def load_master_and_listing(path=MASTER_FILE):
    """read_only で開く(保存しない)。products: pid→dict / listings: list[dict]。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[SH_MASTER]
        p = _header_pos(ws, MASTER_HEAD, SH_MASTER)
        products = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not r[p['pid']]:
                continue
            g = lambda k: r[p[k]] if p[k] < len(r) else None
            products[str(r[p['pid']]).strip()] = {
                'jan': id_str(g('jan')), 'name': str(g('name') or ''), 'type': str(g('type') or ''),
                'cost': g('cost') if isinstance(g('cost'), (int, float)) else None}
        ws = wb[SH_LISTING]
        q = _header_pos(ws, LISTING_HEAD, SH_LISTING)
        listings = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not r[q['lid']]:
                continue
            g = lambda k: r[q[k]] if q[k] < len(r) else None
            listings.append({
                'lid': str(g('lid')).strip(), 'pid': str(g('pid') or '').strip(), 'ch': str(g('ch') or ''),
                'asin': norm(g('asin')), 'asku': norm(g('asku')),
                'pack': g('pack'), 'unit': str(g('unit') or '').strip(), 'proof': str(g('proof') or '')})
    finally:
        wb.close()
    return products, listings


def load_confirmations(path=CONFIRM_FILE):
    """原価確認記録を読む(read_only)。無ければ空。"""
    if not os.path.exists(path):
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SH_CONFIRM not in wb.sheetnames:
            return []
        rows = list(wb[SH_CONFIRM].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    head = [str(h or '').strip() for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if r and r[0] not in (None, '')]


def _ym(v):
    if v is None:
        return ''
    if hasattr(v, 'year'):
        return f'{v.year:04d}-{v.month:02d}'
    m = re.search(r'(\d{4})\D+(\d{1,2})', str(v))
    return f'{int(m.group(1)):04d}-{int(m.group(2)):02d}' if m else ''


def confirmation_for(records, sku, asin, ym):
    """Amazon の出品(SKU または ASIN)に一致し、ym が適用期間内で無効化されていない記録を新しい順に返す。

    適用期間は cost_confirmations.covers と同じ解釈(終了月が空なら単月)。
    確認日だけで無期限に引き継がない(2026-09-12 §7-5)。
    """
    hits = []
    for r in records:
        if str(r.get('チャネル') or '') != 'Amazon' or r.get('無効化日'):
            continue
        if norm(r.get('SKU/ASIN')) not in {k for k in (sku, asin) if k}:
            continue
        a, b = _ym(r.get('適用開始月')), _ym(r.get('適用終了月'))
        if not a:
            continue
        if a <= ym <= (b or a):
            hits.append(r)
    hits.sort(key=lambda r: str(r.get('確認日') or ''), reverse=True)
    return hits


# ─────────────────────────────── 照合 ───────────────────────────────

PACK_PATTERNS = [
    r'(\d+)\s*(?:個|本|袋|箱|缶|パック|点|枚|種)\s*セット',
    r'[×xX✕]\s*(\d+)\s*(?:個|袋|本|箱|缶|パック|セット|\)|）|$|\s)',
    r'\(\s*x\s*(\d+)\s*\)',
    r'(\d+)\s*(?:袋|箱|缶|パック)\s*(?:入り|入)?\s*(?:\)|）|】)',
]


def title_pack_hints(title):
    """商品名から読める入数の候補(参考情報)。**判定には使わない。**

    「60本入 計240本」のような内容量と区別できないため、2〜50の範囲だけ拾い、
    人が推定理由を読むときの手掛かりとして出すに留める。
    """
    found = []
    for pat in PACK_PATTERNS:
        for m in re.finditer(pat, str(title or '')):
            n = int(m.group(1))
            if 2 <= n <= 50 and n not in found:
                found.append(n)
    return found


def link(rec, listings_by_sku, listings_by_asin, csv_skus):
    """CSV1行 → 出品テーブルの出品。戻り値: (出品 or None, 照合方法, 分類 or None, 補足)。

    分類が None なら紐付け成功。
    """
    sku, asin = norm(rec.get('SKU')), norm(rec.get('ASIN'))
    by_sku = listings_by_sku.get(sku, []) if sku else []
    if len(by_sku) == 1:
        return by_sku[0], 'SKU', None, ''
    if len(by_sku) > 1:
        return None, 'SKU', CAT_AMBIG, f'同じSKUの出品が{len(by_sku)}件: ' + \
            ', '.join(f"{c['lid']}({c['pid']})" for c in by_sku)
    by_asin = listings_by_asin.get(asin, []) if asin else []
    if not by_asin:
        return None, '', CAT_NOLINK, ''
    if len(by_asin) > 1:
        return None, 'ASIN', CAT_AMBIG, f'SKUが出品テーブルに無く、同じASINの出品が{len(by_asin)}件: ' + \
            ', '.join(f"{c['lid']}({c['pid']}/SKU {c['asku'] or '空'})" for c in by_asin)
    cand = by_asin[0]
    if cand['asku'] and cand['asku'] != sku and cand['asku'] in csv_skus:
        # その出品のSKUはCSVの別の行として存在する → この行を同じ出品へ寄せると二重になる
        return None, 'ASIN', CAT_AMBIG, (f"同じASINの出品 {cand['lid']} は別SKU {cand['asku']} で"
                                         'CSVにも存在する(同一ASINの別出品の可能性)')
    note = 'SKUは出品テーブルに無くASINで照合(SKU変更の可能性。出品テーブルのSKUは'
    note += (f"{cand['asku']})" if cand['asku'] else '空欄)')
    return cand, 'ASINのみ', None, note


def _is_int_multiple(r):
    n = round(r)
    return n >= 2 and abs(r - n) <= RATIO_TOL * n, n


def classify(pc, mc, pack, unit):
    """プライスター cost と マスター原価を、出品の原価単位を考慮して比べる。

    戻り値: (分類, 期待原価, 差額, 倍率, 推定理由)
      期待原価 … 出品1点あたりに換算したマスター原価(単品原価×入数 / セット原価そのまま / 単位未確認なら素の値)
      倍率     … プライスターcost ÷ マスター標準原価(換算前)。参考値
    ⚠️ どちらが正しいかは判断しない。単位が未確認なら換算もしない(決めつけない)。
    """
    if pc is None or pc <= 0:
        return CAT_NO_PC, None, None, None, 'プライスターのcostが空または0'
    if mc is None:
        return CAT_NO_MC, None, None, None, '商品マスターの標準原価が空'
    n = pack if isinstance(pack, (int, float)) and pack > 0 else None
    if unit == UNIT_SINGLE and n:
        expected, basis = mc * n, f'出品は単品原価×販売入数{int(n)}で換算'
    elif unit == UNIT_SET:
        expected, basis = mc, '出品はセット原価(換算なし)'
    else:
        expected, basis = mc, '原価単位が未確認のため換算せず比較'
    diff = pc - expected
    ratio = round(pc / mc, 4) if mc else None
    if abs(diff) < 0.5:
        return CAT_MATCH, expected, 0, ratio, basis
    unit_known = unit in (UNIT_SINGLE, UNIT_SET) and not (unit == UNIT_SINGLE and not n)
    if not unit_known and ratio:
        ok, k = _is_int_multiple(ratio)
        if ok:
            return CAT_MULTI, expected, diff, ratio, (
                f'{basis}。cost がマスター原価のほぼ{k}倍 → プライスターがセット全体原価('
                f'{k}個分)、マスターが単品原価の可能性')
        ok, k = _is_int_multiple(1 / ratio)
        if ok:
            return CAT_MULTI, expected, diff, ratio, (
                f'{basis}。cost がマスター原価のほぼ1/{k} → マスターがセット原価、'
                'プライスターが単品原価の可能性')
    if abs(diff) <= SMALL_DIFF:
        why = '少額差(10円以下・端数や入力差の可能性)'
    elif expected and 0.85 <= pc / expected <= 1.15:
        why = '±15%以内(仕入値の変動の可能性)'
    else:
        why = '大きな乖離(入力ミス・構成違い・紐付け違いの可能性)'
    return CAT_DIFF, expected, diff, ratio, f'{basis}。{why}'


def build_rows(csv_rows, products, listings, confirmations, ym):
    amz = [l for l in listings if l['ch'] == 'Amazon']
    by_sku, by_asin = {}, {}
    for l in amz:
        if l['asku']:
            by_sku.setdefault(l['asku'], []).append(l)
        if l['asin']:
            by_asin.setdefault(l['asin'], []).append(l)
    # 楽天出品(新形式の管理番号はASINと同値)は紐付け不可の手掛かりとしてだけ使う
    rak_by_asin = {}
    for l in listings:
        if l['ch'] == '楽天' and l['asin']:
            rak_by_asin.setdefault(l['asin'], []).append(l)
    csv_skus = {norm(r.get('SKU')) for r in csv_rows}

    out = []
    for rec in csv_rows:
        sku, asin = norm(rec.get('SKU')), norm(rec.get('ASIN'))
        pc = to_num(rec.get('cost'))
        listing, method, cat, note = link(rec, by_sku, by_asin, csv_skus)
        row = {
            'file': rec['_file'], 'line': rec['_line'],
            'section': 'FBA' if 'FBA' in rec['_file'].upper() else '自己発送',
            'sku': rec.get('SKU', ''), 'asin': rec.get('ASIN', ''), 'title': rec.get('title', ''),
            'stock': to_num(rec.get('number')), 'price': to_num(rec.get('price')), 'pc': pc,
            'method': method, 'lid': '', 'pid': '', 'jan': '', 'name': '', 'ptype': '',
            'mc': None, 'pack': None, 'unit': '', 'expected': None, 'diff': None, 'ratio': None,
            'hints': title_pack_hints(rec.get('title')), 'confirm': '', 'reason': '', 'cat': cat,
        }
        reasons = [note] if note else []
        if cat == CAT_NOLINK:
            reasons.append('SKU・ASINとも出品テーブル(Amazon)に無い')
            row['name'] = '(CSV)' + row['title']     # マスターに無いのでCSVの商品名を仮表示
        if cat == CAT_AMBIG:
            row['name'] = '(CSV)' + row['title']
        if cat == CAT_NOLINK and asin in rak_by_asin:
            reasons.append('楽天出品にのみ同ASINあり: ' + ', '.join(
                f"{l['lid']}({l['pid']})" for l in rak_by_asin[asin]))
        if listing is not None:
            row.update(lid=listing['lid'], pid=listing['pid'], pack=listing['pack'], unit=listing['unit'])
            prod = products.get(listing['pid'])
            if not listing['pid'] or prod is None:
                row['cat'] = CAT_NOLINK
                reasons.append(f"出品 {listing['lid']} の内部管理ID({listing['pid'] or '空'})が商品マスターに無い")
            else:
                row.update(jan=prod['jan'], name=prod['name'], ptype=prod['type'], mc=prod['cost'])
                c, exp, diff, ratio, why = classify(pc, prod['cost'], listing['pack'], listing['unit'])
                row.update(cat=c, expected=exp, diff=diff, ratio=ratio)
                reasons.append(why)
                if c == CAT_MULTI and row['hints']:
                    k = round(ratio) if ratio and ratio >= 1 else None
                    reasons.append('商品名の入数候補 ' + '/'.join(map(str, row['hints'])) +
                                   (' が倍率と一致' if k in row['hints'] else ' (倍率とは不一致)'))
                if c == CAT_DIFF and prod['type'] == 'セット':
                    reasons.append('マスターの商品種別はセット')
        # 原価確認記録(1販売分の原価 = 出品1点あたり)で裏付けがあるか
        hits = confirmation_for(confirmations, sku, asin, ym)
        if hits:
            h = hits[0]
            rc = to_num(h.get('1販売分の原価'))
            same = rc is not None and pc is not None and abs(rc - pc) < 0.5
            row['confirm'] = (f"{h.get('記録ID')}: 1販売分{int(rc) if rc is not None else '?'}円"
                              f"/入数{h.get('販売入数') or '?'}/{h.get('原価単位') or '単位?'}"
                              f"/{_ym(h.get('適用開始月'))}〜{_ym(h.get('適用終了月')) or _ym(h.get('適用開始月'))}"
                              f" → プライスターcostと{'一致' if same else '不一致'}")
            reasons.append('原価確認記録で' + ('裏付けあり' if same else '確認済み原価と食い違い'))
        row['reason'] = ' / '.join(x for x in reasons if x)
        out.append(row)
    out.sort(key=lambda r: (CAT_ORDER.index(r['cat']),
                            -abs(r['diff'] * (r['stock'] or 0)) if isinstance(r['diff'], (int, float)) else 0))
    return out


# ─────────────────────────────── 出力 ───────────────────────────────

COLUMNS = [
    # (見出し, キー, 幅)
    ('分類', 'cat', 22), ('内部管理ID', 'pid', 10), ('出品ID', 'lid', 10), ('SKU', 'sku', 30),
    ('ASIN', 'asin', 12), ('JAN', 'jan', 15), ('商品名', 'name', 40), ('商品マスター原価', 'mc', 10),
    ('プライスターcost', 'pc', 10), ('差額(cost−期待原価)', 'diff', 10), ('倍率(cost÷マスター)', 'ratio', 9),
    ('販売入数', 'pack', 7), ('原価単位', 'unit', 9), ('推定理由', 'reason', 60), ('推奨対応', 'advice', 45),
    ('承認(1を記入)', None, 8), ('判断(マスター更新/プライスター修正/入数登録/保留)', None, 16),
    ('承認者・メモ', None, 20),
    # 以下は判断材料(参考)
    ('期待原価(出品1点換算)', 'expected', 10), ('マスター商品種別', 'ptype', 8), ('照合方法', 'method', 9),
    ('原価確認記録', 'confirm', 40), ('商品名の入数候補(参考)', 'hints', 10), ('在庫数', 'stock', 6),
    ('在庫評価差(差額×在庫)', 'impact', 10), ('販売価格', 'price', 8), ('CSV区分', 'section', 7),
    ('CSV商品名', 'title', 40), ('CSVファイル', 'file', 28), ('CSV行', 'line', 6),
]
ID_KEYS = {'pid', 'lid', 'sku', 'asin', 'jan'}
FILL = {CAT_DIFF: 'FFC7CE', CAT_MULTI: 'FFEB9C', CAT_AMBIG: 'F4B084', CAT_NOLINK: 'D9D9D9',
        CAT_NO_MC: 'FFFF00', CAT_NO_PC: 'DDEBF7', CAT_MATCH: None}


def _cell_value(r, key):
    if key == 'advice':
        return ADVICE[r['cat']]
    if key == 'impact':
        return round(r['diff'] * r['stock']) if isinstance(r['diff'], (int, float)) and r['stock'] else None
    if key == 'hints':
        return '/'.join(map(str, r['hints']))
    v = r.get(key)
    if isinstance(v, float) and v.is_integer() and key not in ('ratio',):
        return int(v)
    return v


def write_report(rows, out_path, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = '概要'
    ws['A1'] = 'プライスターcost 差分レポート(Step A・読み取り専用)'
    ws['A1'].font = Font(bold=True, size=13)
    lines = [
        ('生成日時', meta['generated']), ('判定月(原価確認記録の適用期間)', meta['ym']),
        ('商品マスター', meta['master']), ('原価確認記録', meta['confirm']),
    ] + [('プライスターCSV', f) for f in meta['csv']] + [
        ('', ''),
        ('注意', 'このレポートは商品マスター・出品テーブル・原価確認記録・KPIを一切変更しない。'),
        ('注意', 'プライスターcostは観測値。どちらが正しいかは承認欄で人が判断する(自動上書きしない)。'),
        ('注意', '「一致」の一部はプライスター由来の値をマスターへ還流したもの。独立した裏付けではない。'),
        ('注意', '倍率・商品名の入数候補は参考情報。セットかどうかは出品ごとに確認する(決めつけない)。'),
        ('', ''), ('分類', '件数'),
    ]
    for i, (a, b) in enumerate(lines, start=3):
        ws.cell(i, 1).value, ws.cell(i, 2).value = a, b
    r0 = 3 + len(lines)
    for i, c in enumerate(CAT_ORDER):
        ws.cell(r0 + i, 1).value = c
        ws.cell(r0 + i, 2).value = sum(1 for r in rows if r['cat'] == c)
        if FILL[c]:
            ws.cell(r0 + i, 1).fill = PatternFill('solid', fgColor=FILL[c])
    ws.cell(r0 + len(CAT_ORDER), 1).value = '合計'
    ws.cell(r0 + len(CAT_ORDER), 2).value = len(rows)
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 110

    ws = wb.create_sheet('差分一覧')
    for c, (h, _k, w) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[cell.column_letter].width = w
    for i, r in enumerate(rows, start=2):
        for c, (_h, k, _w) in enumerate(COLUMNS, 1):
            if k is None:
                continue
            cell = ws.cell(i, c, _cell_value(r, k))
            if k in ID_KEYS:
                cell.number_format = '@'     # 識別子は文字列(先頭0・長い数字を崩さない)
        if FILL[r['cat']]:
            ws.cell(i, 1).fill = PatternFill('solid', fgColor=FILL[r['cat']])
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)


# ─────────────────────────────── main ───────────────────────────────

def main(argv=None):
    ap = argparse.ArgumentParser(description='プライスターcost 差分レポート(読み取り専用)')
    ap.add_argument('csv', nargs='+', help='プライスター在庫CSV(自己発送・FBA)')
    ap.add_argument('--month', default=datetime.now().strftime('%Y-%m'),
                    help='原価確認記録の適用期間を判定する月(YYYY-MM)。既定は今月')
    ap.add_argument('--out', help='出力xlsx(既定: SourceData/Output/pricetar_cost_diff_<日時>.xlsx)')
    a = ap.parse_args(argv)
    if not re.fullmatch(r'\d{4}-\d{2}', a.month):
        raise SystemExit('--month は YYYY-MM 形式で指定してください')

    out = a.out or f"{OUTPUT_DIR}/pricetar_cost_diff_{datetime.now():%Y%m%d_%H%M}.xlsx"
    protected = [MASTER_FILE, CONFIRM_FILE]
    if os.path.abspath(out) in {os.path.abspath(p) for p in protected + a.csv}:
        raise SystemExit('出力先が入力ファイルと同じです。別のパスを指定してください')
    before = {p: sha256(p) for p in protected if os.path.exists(p)}

    csv_rows = []
    for p in a.csv:
        csv_rows += read_pricetar_csv(p)
    products, listings = load_master_and_listing()
    confirmations = load_confirmations()
    rows = build_rows(csv_rows, products, listings, confirmations, a.month)

    def stamp(p):
        return f"{os.path.relpath(p, MP.OS_ROOT_STR)}  (更新 {datetime.fromtimestamp(os.path.getmtime(p)):%Y-%m-%d %H:%M})"
    meta = {'generated': f'{datetime.now():%Y-%m-%d %H:%M}', 'ym': a.month,
            'master': stamp(MASTER_FILE),
            'confirm': stamp(CONFIRM_FILE) if os.path.exists(CONFIRM_FILE) else '(ファイル無し)',
            'csv': [stamp(p) + f'  {sum(1 for r in csv_rows if r["_file"] == os.path.basename(p))}行'
                    for p in a.csv]}
    write_report(rows, out, meta)

    # 読み取り専用であることを実行ごとに確かめる(入力の正本が1バイトでも変わっていたら異常終了)
    after = {p: sha256(p) for p in before}
    changed = [p for p in before if before[p] != after[p]]
    if changed:
        raise SystemExit(f'⚠️ 入力ファイルが変化しました(想定外): {changed}')

    print(f'出力: {out}')
    for c in CAT_ORDER:
        print(f'  {c}: {sum(1 for r in rows if r["cat"] == c)}')
    print(f'  合計: {len(rows)}')
    return rows


if __name__ == '__main__':
    main()
